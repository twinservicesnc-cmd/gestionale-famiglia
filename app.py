import io
import base64
import secrets
from urllib.parse import quote
from datetime import timedelta
from html import escape
import hashlib
import hmac
import json
import re
import sqlite3
import uuid
from datetime import date, datetime
from pathlib import Path

import pandas as pd
import plotly.express as px
import streamlit as st
from PIL import Image
from reportlab.lib.utils import ImageReader
from reportlab.platypus import Image as PdfImage, KeepTogether
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


APP_DIR = Path(__file__).resolve().parent
DB_PATH = APP_DIR / "ancillary.db"
NOLEGGIARE_LOGO = APP_DIR / "noleggiare_logo.png"
SEED_PATH = APP_DIR / "dati_iniziali.json"
CONTRACT_SEED_PATH = APP_DIR / "contratti_iniziali.json"
DAMAGE_SEED_PATH = APP_DIR / "danni_iniziali.json"
ANCILLARY_START_DATE = date(2026, 10, 1)
PERMISSION_AREAS = ["ANCILLARY", "CONTRATTI RA", "ADDEBITO DANNI", "CASSA", "EVENTI SPECIALI", "DOCUMENTI", "AMMINISTRAZIONE"]
CASH_IN_TYPES = ["DEPOSITO", "INCASSO", "RETTIFICA POSITIVA"]
CASH_OUT_TYPES = ["RIMBORSO", "RIMESSA", "PRELIEVO", "RETTIFICA NEGATIVA"]

st.set_page_config(
    page_title="Gestionale Noleggi, Ancillary e Danni",
    page_icon=str(APP_DIR / "logo_gestionale.png"),
    layout="wide",
)


def db():
    conn = sqlite3.connect(DB_PATH, check_same_thread=False)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    with db() as conn:
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS rentals (
                id TEXT PRIMARY KEY,
                created_at TEXT NOT NULL,
                ra TEXT NOT NULL,
                start_date TEXT,
                rental_days INTEGER NOT NULL DEFAULT 0,
                source_raw TEXT,
                source_base TEXT,
                source_details TEXT,
                source_year TEXT,
                source_km TEXT,
                operator TEXT,
                vehicle_type TEXT,
                ancillary TEXT,
                ancillary_cost REAL,
                rental_type TEXT,
                notes TEXT
            )
            """
        )
        count = conn.execute("SELECT COUNT(*) FROM rentals").fetchone()[0]
        if count == 0 and SEED_PATH.exists():
            rows = json.loads(SEED_PATH.read_text(encoding="utf-8"))
            conn.executemany(
                """
                INSERT INTO rentals VALUES (
                    :id,:created_at,:ra,:start_date,:rental_days,:source_raw,
                    :source_base,:source_details,:source_year,:source_km,
                    :operator,:vehicle_type,:ancillary,:ancillary_cost,
                    :rental_type,:notes
                )
                """,
                rows,
            )
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS contracts (
                id TEXT PRIMARY KEY,
                import_file TEXT,
                import_period TEXT,
                contract_date TEXT,
                prefix TEXT,
                number INTEGER,
                ra TEXT NOT NULL,
                start_date TEXT,
                end_date TEXT,
                requested_group TEXT,
                assigned_group TEXT,
                source_raw TEXT,
                source_base TEXT,
                source_details TEXT,
                source_year TEXT,
                source_km TEXT,
                client TEXT,
                duration_days INTEGER,
                kpi1_rpd REAL,
                contract_value REAL,
                kpi_target REAL,
                ancillary_value REAL,
                ancillary_rpd REAL,
                operator TEXT,
                UNIQUE(ra, contract_date)
            )
            """
        )
        contracts_schema = conn.execute(
            "SELECT sql FROM sqlite_master WHERE type = 'table' AND name = 'contracts'"
        ).fetchone()[0]
        if "RA TEXT NOT NULL UNIQUE" in upper(contracts_schema):
            conn.execute("ALTER TABLE contracts RENAME TO contracts_legacy_ra_unique")
            conn.execute(
                """
                CREATE TABLE contracts (
                    id TEXT PRIMARY KEY, import_file TEXT, import_period TEXT,
                    contract_date TEXT, prefix TEXT, number INTEGER,
                    ra TEXT NOT NULL, start_date TEXT, end_date TEXT,
                    requested_group TEXT, assigned_group TEXT, source_raw TEXT,
                    source_base TEXT, source_details TEXT, source_year TEXT,
                    source_km TEXT, client TEXT, duration_days INTEGER,
                    kpi1_rpd REAL, contract_value REAL, kpi_target REAL,
                    ancillary_value REAL, ancillary_rpd REAL, operator TEXT,
                    UNIQUE(ra, contract_date)
                )
                """
            )
            conn.execute("INSERT INTO contracts SELECT * FROM contracts_legacy_ra_unique")
            conn.execute("DROP TABLE contracts_legacy_ra_unique")
        contracts_count = conn.execute("SELECT COUNT(*) FROM contracts").fetchone()[0]
        if contracts_count == 0 and CONTRACT_SEED_PATH.exists():
            contracts = json.loads(CONTRACT_SEED_PATH.read_text(encoding="utf-8"))
            conn.executemany(
                """
                INSERT INTO contracts VALUES (
                    :id,:import_file,:import_period,:contract_date,:prefix,:number,
                    :ra,:start_date,:end_date,:requested_group,:assigned_group,
                    :source_raw,:source_base,:source_details,:source_year,:source_km,
                    :client,:duration_days,:kpi1_rpd,:contract_value,:kpi_target,
                    :ancillary_value,:ancillary_rpd,:operator
                )
                """,
                contracts,
            )
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS operator_config (
                name TEXT PRIMARY KEY,
                active INTEGER NOT NULL DEFAULT 1
            )
            """
        )
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS damage_charges (
                id TEXT PRIMARY KEY,
                source_key TEXT NOT NULL UNIQUE,
                submitted_at TEXT,
                control_date TEXT,
                ra TEXT NOT NULL,
                vehicle_category TEXT,
                charge_mode TEXT,
                operator TEXT,
                description TEXT,
                photo_url TEXT,
                amount REAL,
                payment_status TEXT,
                notes TEXT
            )
            """
        )
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS damage_photos (
                id TEXT PRIMARY KEY,
                damage_id TEXT NOT NULL,
                file_name TEXT NOT NULL,
                mime_type TEXT,
                file_data BLOB NOT NULL,
                created_at TEXT NOT NULL
            )
            """
        )
        conn.execute("CREATE INDEX IF NOT EXISTS idx_damage_photos_damage_id ON damage_photos(damage_id)")
        damage_count = conn.execute("SELECT COUNT(*) FROM damage_charges").fetchone()[0]
        if damage_count == 0 and DAMAGE_SEED_PATH.exists():
            damages = json.loads(DAMAGE_SEED_PATH.read_text(encoding="utf-8"))
            conn.executemany(
                """
                INSERT OR IGNORE INTO damage_charges VALUES (
                    :id,:source_key,:submitted_at,:control_date,:ra,
                    :vehicle_category,:charge_mode,:operator,:description,
                    :photo_url,:amount,:payment_status,:notes
                )
                """,
                damages,
            )
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS app_meta (
                key TEXT PRIMARY KEY,
                value TEXT
            )
            """
        )
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS app_users (
                username TEXT PRIMARY KEY,
                password_hash TEXT NOT NULL,
                display_name TEXT,
                operator_name TEXT,
                permissions TEXT NOT NULL DEFAULT '[]',
                active INTEGER NOT NULL DEFAULT 1,
                is_admin INTEGER NOT NULL DEFAULT 0,
                created_at TEXT NOT NULL
            )
            """
        )
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS cash_movements (
                id TEXT PRIMARY KEY,
                source_key TEXT UNIQUE,
                movement_date TEXT NOT NULL,
                ra TEXT,
                movement_type TEXT NOT NULL,
                amount REAL NOT NULL,
                payment_method TEXT NOT NULL DEFAULT 'CONTANTI',
                operator TEXT,
                notes TEXT,
                created_at TEXT NOT NULL
            )
            """
        )
        conn.execute("CREATE INDEX IF NOT EXISTS idx_cash_movements_date ON cash_movements(movement_date)")
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS cash_closings (
                id TEXT PRIMARY KEY,
                closing_date TEXT NOT NULL UNIQUE,
                expected_balance REAL NOT NULL,
                counted_cash REAL NOT NULL,
                checks_total REAL NOT NULL DEFAULT 0,
                difference REAL NOT NULL,
                denominations TEXT NOT NULL DEFAULT '{}',
                operator TEXT,
                notes TEXT,
                created_at TEXT NOT NULL
            )
            """
        )
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS special_events (
                id TEXT PRIMARY KEY,
                title TEXT NOT NULL,
                event_type TEXT,
                start_date TEXT NOT NULL,
                end_date TEXT,
                start_time TEXT,
                end_time TEXT,
                location TEXT,
                responsible TEXT,
                status TEXT NOT NULL DEFAULT 'PROGRAMMATO',
                capacity INTEGER,
                vehicles TEXT,
                cost REAL,
                revenue REAL,
                description TEXT,
                notes TEXT,
                created_by TEXT,
                created_at TEXT NOT NULL
            )
            """
        )
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS event_participants (
                id TEXT PRIMARY KEY,
                event_id TEXT NOT NULL,
                name TEXT NOT NULL,
                role TEXT,
                contact TEXT,
                notes TEXT,
                created_at TEXT NOT NULL
            )
            """
        )
        conn.execute("CREATE INDEX IF NOT EXISTS idx_event_participants_event ON event_participants(event_id)")
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS event_files (
                id TEXT PRIMARY KEY,
                event_id TEXT NOT NULL,
                file_name TEXT NOT NULL,
                mime_type TEXT,
                file_data BLOB NOT NULL,
                created_at TEXT NOT NULL
            )
            """
        )
        conn.execute("CREATE INDEX IF NOT EXISTS idx_event_files_event ON event_files(event_id)")
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS event_vehicles (
                id TEXT PRIMARY KEY,
                event_id TEXT NOT NULL,
                vehicle_group TEXT,
                plate TEXT,
                brand TEXT,
                model TEXT,
                assigned_to TEXT,
                ra TEXT,
                pickup_date TEXT,
                notes TEXT,
                updated_at TEXT NOT NULL
            )
            """
        )
        conn.execute("CREATE INDEX IF NOT EXISTS idx_event_vehicles_event ON event_vehicles(event_id)")
        conn.execute(
            """CREATE TABLE IF NOT EXISTS transport_documents (
                id TEXT PRIMARY KEY, number TEXT NOT NULL UNIQUE, document_date TEXT NOT NULL,
                loading_date TEXT, departure TEXT, destination TEXT, carrier TEXT,
                truck_plate TEXT, station_signature TEXT, driver_signature TEXT,
                delivery_signature TEXT, vehicles_json TEXT NOT NULL DEFAULT '[]',
                created_at TEXT NOT NULL, updated_at TEXT NOT NULL
            )"""
        )
        conn.execute("CREATE INDEX IF NOT EXISTS idx_transport_documents_date ON transport_documents(document_date)")
        conn.execute("""CREATE TABLE IF NOT EXISTS document_assets (
            document_id TEXT NOT NULL, kind TEXT NOT NULL, image_data BLOB NOT NULL,
            signer_name TEXT, signed_at TEXT, PRIMARY KEY(document_id,kind)
        )""")
        conn.execute("""CREATE TABLE IF NOT EXISTS document_sign_requests (
            token_hash TEXT PRIMARY KEY, document_id TEXT NOT NULL,
            role TEXT NOT NULL, expires_at TEXT NOT NULL, signed_at TEXT
        )""")

        cleanup_done = conn.execute(
            "SELECT value FROM app_meta WHERE key = 'ancillary_history_cleared_2026_10_01'"
        ).fetchone()
        if cleanup_done is None:
            conn.execute("DELETE FROM rentals")
            conn.execute(
                "INSERT INTO app_meta (key, value) VALUES (?, ?)",
                ("ancillary_history_cleared_2026_10_01", datetime.now().isoformat(timespec="seconds")),
            )
        operator_count = conn.execute("SELECT COUNT(*) FROM operator_config").fetchone()[0]
        if operator_count == 0:
            existing_operators = conn.execute(
                """
                SELECT DISTINCT TRIM(operator) AS name FROM rentals
                WHERE TRIM(COALESCE(operator, '')) <> ''
                UNION
                SELECT DISTINCT TRIM(operator) AS name FROM contracts
                WHERE TRIM(COALESCE(operator, '')) <> ''
                UNION
                SELECT DISTINCT TRIM(operator) AS name FROM damage_charges
                WHERE TRIM(COALESCE(operator, '')) <> ''
                """
            ).fetchall()
            conn.executemany(
                "INSERT OR IGNORE INTO operator_config (name, active) VALUES (?, 1)",
                [(upper(row[0]),) for row in existing_operators if upper(row[0])],
            )


def clean(value):
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass
    return re.sub(r"\s+", " ", str(value or "").strip())


def upper(value):
    return clean(value).upper()


def configured_operators(active_only=True):
    query = "SELECT name, active FROM operator_config"
    if active_only:
        query += " WHERE active = 1"
    query += " ORDER BY name"
    with db() as conn:
        return [dict(row) for row in conn.execute(query).fetchall()]


def password_hash(password, salt=None):
    salt = salt or uuid.uuid4().hex
    digest = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt.encode("utf-8"), 200_000).hex()
    return f"pbkdf2_sha256${salt}${digest}"


def password_matches(password, stored_hash):
    try:
        algorithm, salt, expected = stored_hash.split("$", 2)
        if algorithm != "pbkdf2_sha256":
            return False
        actual = password_hash(password, salt).split("$", 2)[2]
        return hmac.compare_digest(actual, expected)
    except (AttributeError, ValueError):
        return False


def current_user():
    return st.session_state.get("current_user", {})


def current_operator():
    return upper(current_user().get("operator_name", ""))


def user_is_admin():
    return bool(current_user().get("is_admin"))


def allowed_areas():
    if user_is_admin():
        return set(PERMISSION_AREAS)
    return set(current_user().get("permissions", []))


def normalize_ra(value):
    text = upper(value).replace(".", "").replace(" ", "").replace("_", "-")
    match = re.search(r"(?:TOR)?-?(\d+)", text)
    return f"TOR-{match.group(1)}" if match else text


def parse_source(value):
    text = upper(value)
    for wrong, right in {
        "TRAVELJGSAW": "TRAVELJIGSAW",
        "WALK IN": "WALK-IN",
        "ILLIMITATI": "UNLIMITED",
    }.items():
        text = text.replace(wrong, right)
    year_match = re.search(r"\b(20\d{2})\b", text)
    km_match = re.search(r"\b(\d{2,5})\s*KM\b", text)
    known = [
        "OFFICIAL BOOKING WEB", "TRAVELJIGSAW", "TRAVELDRIVE",
        "VIPCARS", "DOYOUSPAIN", "WALK-IN", "RECEPTION",
        "POSTEGO", "POSTE", "WEB",
    ]
    base = next((item for item in known if text.startswith(item)), "")
    if not base:
        marker = re.search(r"\s+(?:AUTO|VAN|CAR)\b", text)
        base = text[: marker.start()].strip() if marker else text
    details = text[len(base):].strip(" -")
    return {
        "source_raw": text,
        "source_base": base,
        "source_details": details,
        "source_year": year_match.group(1) if year_match else "",
        "source_km": km_match.group(1) if km_match else "",
    }


def classify_contract_channel(source_raw):
    text = upper(source_raw)
    if "REPLACEMENT" in text or "WARRANTY" in text:
        return "REPLACEMENT"
    if "WALK-IN" in text or "WALK IN" in text or text.startswith("RECEPTION"):
        return "WALK-IN"
    if "OFFICIAL BOOKING WEB" in text or re.search(r"\b(?:WEB|CNP)\b", text):
        return "WEB/CNP"
    broker_markers = (
        "TRAVELJIGSAW", "TRAVELDRIVE", "DOYOUSPAIN", "VIPCARS",
        "BOOKING GROUP", "TINOLEGGIO", "XML", "POA",
    )
    if any(marker in text for marker in broker_markers):
        return "BROKER"
    return "CORPORATE"


def normalize_rental_type(value):
    text = upper(value).replace("WEB-CNP-ETC", "WEB - CNP ETC").replace("CORPARATE", "CORPORATE")
    parts = []
    for item in text.split(","):
        item = clean(item)
        if item and item not in parts:
            parts.append(item)
    return ", ".join(parts)


def save_record(record):
    source = parse_source(record["source_raw"])
    payload = {
        "id": record.get("id") or str(uuid.uuid4()),
        "created_at": record.get("created_at") or datetime.now().isoformat(timespec="seconds"),
        "ra": normalize_ra(record["ra"]),
        "start_date": str(record.get("start_date") or ""),
        "rental_days": int(record.get("rental_days") or 0),
        **source,
        "operator": upper(record.get("operator")),
        "vehicle_type": upper(record.get("vehicle_type")),
        "ancillary": upper(record.get("ancillary")).replace("PRESNETE", "PRESENTE"),
        "ancillary_cost": record.get("ancillary_cost"),
        "rental_type": normalize_rental_type(record.get("rental_type")),
        "notes": clean(record.get("notes")),
    }
    with db() as conn:
        conn.execute(
            """
            INSERT INTO rentals VALUES (
                :id,:created_at,:ra,:start_date,:rental_days,:source_raw,
                :source_base,:source_details,:source_year,:source_km,
                :operator,:vehicle_type,:ancillary,:ancillary_cost,
                :rental_type,:notes
            )
            ON CONFLICT(id) DO UPDATE SET
                created_at=:created_at, ra=:ra, start_date=:start_date,
                rental_days=:rental_days, source_raw=:source_raw,
                source_base=:source_base, source_details=:source_details,
                source_year=:source_year, source_km=:source_km,
                operator=:operator, vehicle_type=:vehicle_type,
                ancillary=:ancillary, ancillary_cost=:ancillary_cost,
                rental_type=:rental_type, notes=:notes
            """,
            payload,
        )


def load_data():
    with db() as conn:
        frame = pd.read_sql_query("SELECT * FROM rentals ORDER BY start_date DESC, created_at DESC", conn)
    if frame.empty:
        return frame
    frame["start_date"] = pd.to_datetime(frame["start_date"], errors="coerce")
    frame["rental_days"] = pd.to_numeric(frame["rental_days"], errors="coerce").fillna(0).astype(int)
    frame["ancillary_cost"] = pd.to_numeric(frame["ancillary_cost"], errors="coerce")
    frame["rpd"] = frame["ancillary_cost"].div(frame["rental_days"].replace(0, pd.NA))
    frame["has_ancillary"] = frame["ancillary"].fillna("").str.strip().ne("") & frame["ancillary_cost"].fillna(0).gt(0)
    return frame


def classify_damage(description):
    text = upper(description)
    categories = [
        ("PARAURTI", "PARAURTI"), ("PARABREZZA", "PARABREZZA"),
        ("PARAFANGO", "PARAFANGO"), ("PORTA", "PORTA/PORTIERA"),
        ("PORTIER", "PORTA/PORTIERA"), ("FIANC", "FIANCATA"),
        ("SPECCHI", "SPECCHIETTO"), ("CERCH", "CERCHIO/PNEUMATICO"),
        ("PNEUM", "CERCHIO/PNEUMATICO"), ("FARO", "FARO/FANALE"),
        ("FANAL", "FARO/FANALE"), ("TETTO", "TETTO"),
        ("INTERN", "INTERNI"), ("BOLLO", "AMMACCATURA/BOLLO"),
        ("AMMAC", "AMMACCATURA/BOLLO"),
    ]
    return next((label for marker, label in categories if marker in text), "ALTRO/DA CLASSIFICARE")


def save_damage(record):
    record_id = record.get("id") or str(uuid.uuid4())
    submitted_at = record.get("submitted_at") or datetime.now().isoformat(timespec="seconds")
    source_key = clean(record.get("source_key")) or record_id
    payload = {
        "id": record_id,
        "source_key": source_key,
        "submitted_at": submitted_at,
        "control_date": clean(record.get("control_date")),
        "ra": normalize_ra(record.get("ra")),
        "vehicle_category": upper(record.get("vehicle_category")),
        "charge_mode": upper(record.get("charge_mode")),
        "operator": upper(record.get("operator")),
        "description": upper(record.get("description")),
        "photo_url": clean(record.get("photo_url")),
        "amount": record.get("amount"),
        "payment_status": upper(record.get("payment_status")) or "DA VERIFICARE",
        "notes": clean(record.get("notes")),
    }
    with db() as conn:
        conn.execute(
            """
            INSERT INTO damage_charges VALUES (
                :id,:source_key,:submitted_at,:control_date,:ra,
                :vehicle_category,:charge_mode,:operator,:description,
                :photo_url,:amount,:payment_status,:notes
            )
            ON CONFLICT(source_key) DO UPDATE SET
                submitted_at=:submitted_at, control_date=:control_date, ra=:ra,
                vehicle_category=:vehicle_category, charge_mode=:charge_mode,
                operator=:operator, description=:description, photo_url=:photo_url,
                amount=:amount, payment_status=:payment_status, notes=:notes
            """,
            payload,
        )
    return record_id


def save_damage_photos(damage_id, uploaded_files):
    rows = []
    for uploaded in uploaded_files or []:
        file_data = uploaded.getvalue()
        if not file_data:
            continue
        rows.append((
            str(uuid.uuid4()), damage_id, clean(uploaded.name) or "foto_danno.jpg",
            clean(uploaded.type) or "image/jpeg", sqlite3.Binary(file_data),
            datetime.now().isoformat(timespec="seconds"),
        ))
    if rows:
        with db() as conn:
            conn.executemany(
                """
                INSERT INTO damage_photos
                (id, damage_id, file_name, mime_type, file_data, created_at)
                VALUES (?, ?, ?, ?, ?, ?)
                """,
                rows,
            )
    return len(rows)


def damage_photos(damage_id):
    with db() as conn:
        return [dict(row) for row in conn.execute(
            "SELECT * FROM damage_photos WHERE damage_id = ? ORDER BY created_at", (damage_id,)
        ).fetchall()]


def load_damages():
    with db() as conn:
        frame = pd.read_sql_query(
            "SELECT * FROM damage_charges ORDER BY control_date DESC, submitted_at DESC", conn
        )
    if frame.empty:
        return frame
    frame["submitted_at"] = pd.to_datetime(frame["submitted_at"], errors="coerce")
    frame["control_date"] = pd.to_datetime(frame["control_date"], errors="coerce")
    frame["amount"] = pd.to_numeric(frame["amount"], errors="coerce")
    frame["damage_category"] = frame["description"].map(classify_damage)
    with db() as conn:
        photo_counts = pd.read_sql_query(
            "SELECT damage_id AS id, COUNT(*) AS photo_count FROM damage_photos GROUP BY damage_id", conn
        )
    if photo_counts.empty:
        frame["photo_count"] = 0
    else:
        frame = frame.merge(photo_counts, on="id", how="left")
        frame["photo_count"] = frame["photo_count"].fillna(0).astype(int)
    frame["has_photo"] = frame["photo_url"].fillna("").str.strip().ne("") | frame["photo_count"].gt(0)
    return frame


def cash_effect(movement_type, amount):
    value = abs(float(amount or 0))
    return value if upper(movement_type) in CASH_IN_TYPES else -value


def save_cash_movement(record):
    movement_id = record.get("id") or str(uuid.uuid4())
    payload = {
        "id": movement_id,
        "source_key": clean(record.get("source_key")) or movement_id,
        "movement_date": clean(record.get("movement_date")) or date.today().isoformat(),
        "ra": normalize_ra(record.get("ra")) if clean(record.get("ra")) else "",
        "movement_type": upper(record.get("movement_type")),
        "amount": abs(float(record.get("amount") or 0)),
        "payment_method": upper(record.get("payment_method")) or "CONTANTI",
        "operator": upper(record.get("operator")),
        "notes": clean(record.get("notes")),
        "created_at": record.get("created_at") or datetime.now().isoformat(timespec="seconds"),
    }
    with db() as conn:
        conn.execute(
            """
            INSERT INTO cash_movements
            (id, source_key, movement_date, ra, movement_type, amount, payment_method, operator, notes, created_at)
            VALUES (:id, :source_key, :movement_date, :ra, :movement_type, :amount, :payment_method, :operator, :notes, :created_at)
            ON CONFLICT(source_key) DO UPDATE SET
                movement_date=:movement_date, ra=:ra, movement_type=:movement_type,
                amount=:amount, payment_method=:payment_method, operator=:operator,
                notes=:notes, created_at=:created_at
            """,
            payload,
        )
    return movement_id


def load_cash_movements():
    with db() as conn:
        frame = pd.read_sql_query(
            "SELECT * FROM cash_movements ORDER BY movement_date DESC, created_at DESC", conn
        )
    if frame.empty:
        return frame
    frame["movement_date"] = pd.to_datetime(frame["movement_date"], errors="coerce")
    frame["amount"] = pd.to_numeric(frame["amount"], errors="coerce").fillna(0)
    frame["effect"] = frame.apply(lambda row: cash_effect(row["movement_type"], row["amount"]), axis=1)
    return frame


def cash_balance(frame=None):
    frame = load_cash_movements() if frame is None else frame
    return 0.0 if frame.empty else float(frame["effect"].sum())


def load_cash_closings():
    with db() as conn:
        frame = pd.read_sql_query("SELECT * FROM cash_closings ORDER BY closing_date DESC", conn)
    if frame.empty:
        return frame
    frame["closing_date"] = pd.to_datetime(frame["closing_date"], errors="coerce")
    for column in ["expected_balance", "counted_cash", "checks_total", "difference"]:
        frame[column] = pd.to_numeric(frame[column], errors="coerce").fillna(0)
    return frame


def save_special_event(record):
    event_id = record.get("id") or str(uuid.uuid4())
    payload = {
        "id": event_id, "title": clean(record.get("title")),
        "event_type": upper(record.get("event_type")),
        "start_date": clean(record.get("start_date")), "end_date": clean(record.get("end_date")),
        "start_time": clean(record.get("start_time")), "end_time": clean(record.get("end_time")),
        "location": clean(record.get("location")), "responsible": upper(record.get("responsible")),
        "status": upper(record.get("status")) or "PROGRAMMATO",
        "capacity": int(record.get("capacity") or 0), "vehicles": clean(record.get("vehicles")),
        "cost": float(record.get("cost") or 0), "revenue": float(record.get("revenue") or 0),
        "description": clean(record.get("description")), "notes": clean(record.get("notes")),
        "created_by": clean(record.get("created_by")) or current_user().get("username", ""),
        "created_at": record.get("created_at") or datetime.now().isoformat(timespec="seconds"),
    }
    with db() as conn:
        conn.execute(
            """
            INSERT INTO special_events VALUES (
                :id,:title,:event_type,:start_date,:end_date,:start_time,:end_time,
                :location,:responsible,:status,:capacity,:vehicles,:cost,:revenue,
                :description,:notes,:created_by,:created_at
            )
            ON CONFLICT(id) DO UPDATE SET
                title=:title,event_type=:event_type,start_date=:start_date,end_date=:end_date,
                start_time=:start_time,end_time=:end_time,location=:location,
                responsible=:responsible,status=:status,capacity=:capacity,vehicles=:vehicles,
                cost=:cost,revenue=:revenue,description=:description,notes=:notes,
                created_by=:created_by,created_at=:created_at
            """, payload,
        )
    return event_id


def load_special_events():
    with db() as conn:
        frame = pd.read_sql_query("SELECT * FROM special_events ORDER BY start_date DESC, start_time DESC", conn)
    if frame.empty:
        return frame
    frame["start_date"] = pd.to_datetime(frame["start_date"], errors="coerce")
    frame["end_date"] = pd.to_datetime(frame["end_date"], errors="coerce")
    for column in ["capacity", "cost", "revenue"]:
        frame[column] = pd.to_numeric(frame[column], errors="coerce").fillna(0)
    frame["margin"] = frame["revenue"] - frame["cost"]
    with db() as conn:
        counts = pd.read_sql_query("SELECT event_id AS id, COUNT(*) AS participants FROM event_participants GROUP BY event_id", conn)
        files = pd.read_sql_query("SELECT event_id AS id, COUNT(*) AS files FROM event_files GROUP BY event_id", conn)
        vehicles = pd.read_sql_query("SELECT event_id AS id, COUNT(*) AS assigned_vehicles FROM event_vehicles GROUP BY event_id", conn)
    frame = frame.merge(counts, on="id", how="left") if not counts.empty else frame.assign(participants=0)
    frame = frame.merge(files, on="id", how="left") if not files.empty else frame.assign(files=0)
    frame = frame.merge(vehicles, on="id", how="left") if not vehicles.empty else frame.assign(assigned_vehicles=0)
    frame[["participants", "files", "assigned_vehicles"]] = frame[["participants", "files", "assigned_vehicles"]].fillna(0).astype(int)
    return frame


def event_options(frame):
    if frame.empty:
        return {}, []
    records = {row.id: row for row in frame.itertuples()}
    return records, list(records)


def load_contracts():
    with db() as conn:
        frame = pd.read_sql_query("SELECT * FROM contracts ORDER BY contract_date DESC, ra DESC", conn)
    if frame.empty:
        return frame
    for column in ["contract_date", "start_date", "end_date"]:
        frame[column] = pd.to_datetime(frame[column], errors="coerce")
    for column in ["duration_days", "kpi1_rpd", "contract_value", "kpi_target", "ancillary_value", "ancillary_rpd"]:
        frame[column] = pd.to_numeric(frame[column], errors="coerce")
    frame["has_ancillary_ra"] = frame["ancillary_value"].fillna(0).gt(0)
    frame["group_changed"] = frame["requested_group"].fillna("").ne(frame["assigned_group"].fillna(""))
    frame["contract_vehicle"] = frame["assigned_group"].fillna("").map(
        lambda value: "VAN" if upper(value).startswith("Z") else "CAR"
    )
    frame["rental_term"] = frame.apply(
        lambda row: "MENSILE"
        if "MENSILE" in upper(row["source_raw"]) or (row["duration_days"] or 0) >= 28
        else "GIORNALIERO",
        axis=1,
    )
    frame["rental_channel"] = frame["source_raw"].map(classify_contract_channel)
    return frame


def _find_header_row(raw):
    for index in range(min(25, len(raw))):
        values = {clean(value) for value in raw.iloc[index].dropna().tolist()}
        if {"Data contratto", "Prefisso", "Numero"}.issubset(values):
            return index
    return None


def detect_excel_type(file_bytes):
    book = pd.ExcelFile(io.BytesIO(file_bytes))
    for sheet in book.sheet_names:
        raw = pd.read_excel(io.BytesIO(file_bytes), sheet_name=sheet, header=None, nrows=25, dtype=object)
        if _find_header_row(raw) is not None:
            return "CONTRATTI_RA"
    first = pd.read_excel(io.BytesIO(file_bytes), sheet_name=0, nrows=5, dtype=object)
    damage_columns = {
        "Informazioni cronologiche", "Numero RA (Rental Agreement)",
        "Categoria Veicolo", "Modalità di addebito", "Operatore responsabile",
    }
    if damage_columns.issubset(first.columns):
        return "ADDEBITO_DANNI"
    if {"RA (Rental Agreement)", "DATA INIZIO NOLEGGIO", "GIORNI NOLEGGIO", "FONTE"}.issubset(first.columns):
        return "ANCILLARY"
    return "SCONOSCIUTO"


def import_contract_workbook(file_bytes, file_name):
    book = pd.ExcelFile(io.BytesIO(file_bytes))
    parsed = {}
    for sheet in book.sheet_names:
        raw = pd.read_excel(io.BytesIO(file_bytes), sheet_name=sheet, header=None, nrows=25, dtype=object)
        header = _find_header_row(raw)
        if header is None:
            continue
        frame = pd.read_excel(io.BytesIO(file_bytes), sheet_name=sheet, header=header, dtype=object)
        frame = frame[frame["Numero"].notna()].copy()
        if not frame.empty:
            parsed[sheet] = frame
    if not parsed:
        raise ValueError("Nessun foglio RA riconosciuto.")
    site_name = max(parsed, key=lambda name: len(parsed[name]))
    site = parsed[site_name].copy()

    def ra_from_row(row):
        number = pd.to_numeric(row.get("Numero"), errors="coerce")
        if pd.isna(number):
            return ""
        return f"{upper(row.get('Prefisso'))}-{int(number)}"

    operator_map = {}
    for sheet, frame in parsed.items():
        if sheet == site_name or sheet.upper() == "TOTALE":
            continue
        for _, row in frame.iterrows():
            operator_map[ra_from_row(row)] = clean(sheet)

    contract_dates = pd.to_datetime(site["Data contratto"], errors="coerce")
    valid_dates = contract_dates.dropna()
    period = valid_dates.min().strftime("%Y-%m") if not valid_dates.empty else ""
    imported = 0
    with db() as conn:
        for _, row in site.iterrows():
            ra = ra_from_row(row)
            if not ra:
                continue
            source = parse_source(row.get("Fonte commissione"))
            number = pd.to_numeric(row.get("Numero"), errors="coerce")

            def number_value(column):
                value = pd.to_numeric(row.get(column), errors="coerce")
                return None if pd.isna(value) else float(value)

            def date_value(column):
                value = pd.to_datetime(row.get(column), errors="coerce")
                return "" if pd.isna(value) else value.date().isoformat()

            duration = number_value("Durata (gg)") or 0
            payload = {
                "id": str(uuid.uuid4()), "import_file": file_name, "import_period": period,
                "contract_date": date_value("Data contratto"), "prefix": upper(row.get("Prefisso")),
                "number": int(number), "ra": ra, "start_date": date_value("Data inizio contratto"),
                "end_date": date_value("Data fine contratto"), "requested_group": upper(row.get("Gruppo richiesto")),
                "assigned_group": upper(row.get("Gruppo assegnato")), **source,
                "client": clean(row.get("Cliente")), "duration_days": int(round(duration)),
                "kpi1_rpd": number_value("KPI 1"), "contract_value": number_value("Valore contratto 1"),
                "kpi_target": number_value("KPI 1 Ob."), "ancillary_value": number_value("Valore contratto 2"),
                "ancillary_rpd": number_value("RpD 2"), "operator": operator_map.get(ra, "NON ASSEGNATO"),
            }
            conn.execute(
                """
                INSERT INTO contracts VALUES (
                    :id,:import_file,:import_period,:contract_date,:prefix,:number,
                    :ra,:start_date,:end_date,:requested_group,:assigned_group,
                    :source_raw,:source_base,:source_details,:source_year,:source_km,
                    :client,:duration_days,:kpi1_rpd,:contract_value,:kpi_target,
                    :ancillary_value,:ancillary_rpd,:operator
                )
                ON CONFLICT(ra, contract_date) DO UPDATE SET
                    import_file=:import_file, import_period=:import_period,
                    contract_date=:contract_date, prefix=:prefix, number=:number,
                    start_date=:start_date, end_date=:end_date,
                    requested_group=:requested_group, assigned_group=:assigned_group,
                    source_raw=:source_raw, source_base=:source_base,
                    source_details=:source_details, source_year=:source_year,
                    source_km=:source_km, client=:client, duration_days=:duration_days,
                    kpi1_rpd=:kpi1_rpd, contract_value=:contract_value,
                    kpi_target=:kpi_target, ancillary_value=:ancillary_value,
                    ancillary_rpd=:ancillary_rpd, operator=:operator
                """,
                payload,
            )
            imported += 1
    return imported, site_name, period


def import_damage_workbook(file_bytes, file_name):
    book = pd.ExcelFile(io.BytesIO(file_bytes))
    imported = 0
    skipped = 0
    for sheet in book.sheet_names:
        frame = pd.read_excel(io.BytesIO(file_bytes), sheet_name=sheet, dtype=object)
        required = {
            "Informazioni cronologiche", "Numero RA (Rental Agreement)",
            "Categoria Veicolo", "Modalità di addebito", "Operatore responsabile",
        }
        if not required.issubset(frame.columns):
            continue
        for _, row in frame.iterrows():
            ra = normalize_ra(row.get("Numero RA (Rental Agreement)"))
            if not ra:
                skipped += 1
                continue
            submitted = pd.to_datetime(row.get("Informazioni cronologiche"), errors="coerce")
            control = pd.to_datetime(row.get("Data del controllo"), errors="coerce")
            submitted_text = "" if pd.isna(submitted) else submitted.isoformat()
            source_key = f"{submitted_text}|{ra}|{sheet}"
            save_damage({
                "source_key": source_key,
                "submitted_at": submitted_text,
                "control_date": "" if pd.isna(control) else control.date().isoformat(),
                "ra": ra,
                "vehicle_category": row.get("Categoria Veicolo"),
                "charge_mode": row.get("Modalità di addebito"),
                "operator": row.get("Operatore responsabile"),
                "description": row.get("Descrizione sintetica del danno"),
                "photo_url": row.get("Caricamento foto danno (opzionale)"),
                "amount": None,
                "payment_status": "DA VERIFICARE",
                "notes": f"Importato da {file_name} - foglio {sheet}",
            })
            imported += 1
    return imported, skipped


def to_excel(frame):
    output = io.BytesIO()
    export = frame.copy()
    export["DATA INIZIO NOLEGGIO"] = export["start_date"].dt.date
    export = export.rename(columns={
        "ra": "RA", "rental_days": "GIORNI NOLEGGIO", "source_raw": "FONTE COMPLETA",
        "source_base": "FONTE PRINCIPALE", "source_details": "DETTAGLIO FONTE",
        "source_year": "ANNO FONTE", "source_km": "KM FONTE", "operator": "OPERATORE",
        "vehicle_type": "TIPO VEICOLO", "ancillary": "ANCILLARY",
        "ancillary_cost": "COSTO ANCILLARY IVA ESCLUSA", "rpd": "RPD",
        "rental_type": "TIPO NOLEGGIO", "notes": "NOTE",
    })
    cols = [
        "RA", "DATA INIZIO NOLEGGIO", "GIORNI NOLEGGIO", "FONTE COMPLETA",
        "FONTE PRINCIPALE", "DETTAGLIO FONTE", "ANNO FONTE", "KM FONTE",
        "OPERATORE", "TIPO VEICOLO", "ANCILLARY", "COSTO ANCILLARY IVA ESCLUSA",
        "RPD", "TIPO NOLEGGIO", "NOTE",
    ]
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        export[cols].to_excel(writer, index=False, sheet_name="Dati ancillary")
    return output.getvalue()


def table_to_excel(frame, sheet_name="Dati"):
    output = io.BytesIO()
    export = frame.copy()
    for column in export.columns:
        if pd.api.types.is_datetime64_any_dtype(export[column]):
            export[column] = export[column].dt.date
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        export.to_excel(writer, index=False, sheet_name=sheet_name[:31])
    return output.getvalue()


def next_ddt_number(conn, year):
    rows = conn.execute("SELECT number FROM transport_documents WHERE document_date LIKE ?", (f"{year}-%",)).fetchall()
    used = [int(match.group(1)) for row in rows if (match := re.fullmatch(rf"{year}/(\d+)", row["number"] or ""))]
    return f"{year}/{max(used, default=0) + 1:04d}"


DDT_SIGNATURE_ROLES = {"driver": "Firma autista", "station": "Firma operatore stazione",
                       "delivery": "Firma operatore stazione consegna / piazzale"}
DDT_ASSET_KINDS = {"stamp": "Timbro", **DDT_SIGNATURE_ROLES}


def ddt_assets(document_id):
    with db() as conn:
        return {row["kind"]: dict(row) for row in conn.execute(
            "SELECT * FROM document_assets WHERE document_id = ?", (document_id,))}


def image_for_pdf(raw, max_width, max_height):
    with Image.open(io.BytesIO(raw)) as original:
        img = original.convert("RGBA")
        img.thumbnail((1200, 500))
        background = Image.new("RGB", img.size, "white")
        background.paste(img, mask=img.getchannel("A"))
        buffer = io.BytesIO()
        background.save(buffer, format="PNG")
    buffer.seek(0)
    width, height = background.size
    scale = min(max_width / width, max_height / height)
    return PdfImage(buffer, width=width*scale, height=height*scale)


def ddt_pdf(row):
    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=A4, leftMargin=14*mm, rightMargin=14*mm,
                            topMargin=14*mm, bottomMargin=14*mm)
    styles = getSampleStyleSheet()
    assets = ddt_assets(row["id"])
    def para(value):
        return Paragraph(escape(str(value or "")), styles["Normal"])
    def field(label, value):
        return [para(label), para(value)]
    header = Table([[PdfImage(str(NOLEGGIARE_LOGO), width=50*mm, height=15*mm),
                     Paragraph("DOCUMENTO DI TRASPORTO", styles["Heading1"])]],
                   colWidths=[62*mm, 120*mm], hAlign="LEFT")
    header.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                                ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))
    story = [header, Spacer(1, 5*mm),
             para(f"Bolla n° {row['number']}    ·    Data {row['document_date']}"), Spacer(1, 5*mm)]
    details = [field("Data carico", row["loading_date"]), field("Stazione di partenza", row["departure"]),
               field("Destinazione", row["destination"]), field("Vettore / Autisti", row["carrier"]),
               field("Targa bisarca", row["truck_plate"])]
    table = Table(details, colWidths=[49*mm, 133*mm], hAlign="LEFT")
    table.setStyle(TableStyle([("GRID",(0,0),(-1,-1),0.4,colors.grey),
                               ("VALIGN",(0,0),(-1,-1),"TOP"),
                               ("BACKGROUND",(0,0),(0,-1),colors.HexColor("#EFF3F8")),
                               ("PADDING",(0,0),(-1,-1),6)]))
    story += [table, Spacer(1, 7*mm)]
    vehicles = json.loads(row["vehicles_json"] or "[]")
    headings = ["n°", "Marca", "Modello", "Targa", "Materiale a bordo / Note"]
    entries = [[para(h) for h in headings]]
    entries.extend([[para(i), para(v.get("marca")), para(v.get("modello")),
                     para(v.get("targa")), para(v.get("note"))] for i,v in enumerate(vehicles, 1)])
    entries.extend([[para("") for _ in headings] for _ in range(max(0, 10-len(vehicles)))])
    vehicles_table = Table(entries, colWidths=[12*mm, 31*mm, 35*mm, 27*mm, 77*mm], repeatRows=1, hAlign="LEFT")
    vehicles_table.setStyle(TableStyle([("GRID",(0,0),(-1,-1),0.4,colors.grey),
                                         ("BACKGROUND",(0,0),(-1,0),colors.HexColor("#DCE6F1")),
                                         ("VALIGN",(0,0),(-1,-1),"TOP"),
                                         ("PADDING",(0,0),(-1,-1),5)]))
    story += [vehicles_table, Spacer(1, 9*mm)]
    def signature_content(kind, label, width=62*mm):
        asset = assets.get(kind)
        content = [Paragraph(f"<b>{escape(label)}</b>", styles["Normal"]), Spacer(1, 3*mm)]
        if asset:
            content.append(image_for_pdf(asset["image_data"], width, 18*mm))
            content.append(para(f"{asset['signer_name'] or ''} · {asset['signed_at'] or ''}"))
        else:
            name = row[f"{kind}_signature"]
            content.append(para(name or "________________________________"))
        return content

    # Schema del modello: operatore di partenza e timbro a sinistra;
    # autista in alto a destra, operatore di consegna in basso a destra.
    station_block = [Paragraph("<b>Timbro e firma operatore stazione</b>", styles["Normal"]), Spacer(1, 4*mm)]
    if "stamp" in assets:
        station_block += [image_for_pdf(assets["stamp"]["image_data"], 65*mm, 28*mm), Spacer(1, 3*mm)]
    else:
        station_block += [para("Timbro: ______________________________"), Spacer(1, 5*mm)]
    station_block += signature_content("station", "Firma operatore stazione", 65*mm)[2:]
    signatures = Table([[station_block, signature_content("driver", "Firma autista")],
                        ["", signature_content("delivery", "Firma operatore stazione consegna / piazzale")]],
                       colWidths=[91*mm, 91*mm], rowHeights=[39*mm, 39*mm], hAlign="LEFT")
    signatures.setStyle(TableStyle([("SPAN", (0, 0), (0, 1)),
                                    ("BOX", (0, 0), (0, 1), 0.7, colors.grey),
                                    ("BOX", (1, 0), (1, 0), 0.7, colors.grey),
                                    ("BOX", (1, 1), (1, 1), 0.7, colors.grey),
                                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                    ("LEFTPADDING", (0, 0), (-1, -1), 6),
                                    ("TOPPADDING", (0, 0), (-1, -1), 7)]))
    story.append(KeepTogether(signatures))
    doc.build(story)
    return output.getvalue()


def ddt_excel(row):
    details = {"Numero": row["number"], "Data documento": row["document_date"],
               "Data carico": row["loading_date"], "Partenza": row["departure"],
               "Destinazione": row["destination"], "Vettore / Autisti": row["carrier"],
               "Targa bisarca": row["truck_plate"]}
    vehicles = json.loads(row["vehicles_json"] or "[]")
    assets = ddt_assets(row["id"])
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        pd.DataFrame(list(details.items()), columns=["Campo", "Valore"]).to_excel(writer, index=False, sheet_name="DDT")
        pd.DataFrame([{"N°": i, "Marca": v.get("marca"), "Modello": v.get("modello"),
                       "Targa": v.get("targa"), "Materiale e note": v.get("note")}
                      for i, v in enumerate(vehicles, 1)],
                     columns=["N°", "Marca", "Modello", "Targa", "Materiale e note"]).to_excel(
                         writer, index=False, sheet_name="Veicoli")
        pd.DataFrame([{"Ruolo": label, "Firmatario": assets.get(role, {}).get("signer_name", ""),
                       "Data firma": assets.get(role, {}).get("signed_at", "")}
                      for role, label in DDT_SIGNATURE_ROLES.items()]).to_excel(
                          writer, index=False, sheet_name="Firme")
        from openpyxl.drawing.image import Image as ExcelImage
        from openpyxl.utils import get_column_letter
        sheet = writer.sheets["DDT"]
        logo = ExcelImage(str(NOLEGGIARE_LOGO))
        logo.width, logo.height = 168, 50
        sheet.add_image(logo, "D1")
        sheet.column_dimensions["D"].width = 27
        sheet.column_dimensions["A"].width = 25
        sheet.column_dimensions["B"].width = 55
        for index, (kind, label) in enumerate(DDT_ASSET_KINDS.items(), start=12):
            if kind in assets:
                sheet.cell(index, 1, label)
                img = ExcelImage(io.BytesIO(assets[kind]["image_data"]))
                img.width, img.height = 160, 55
                sheet.add_image(img, f"B{index}")
                sheet.row_dimensions[index].height = 46
        vehicle_sheet = writer.sheets["Veicoli"]
        for col, width in enumerate([8, 22, 24, 20, 50], 1):
            vehicle_sheet.column_dimensions[get_column_letter(col)].width = width
    return output.getvalue()


def validate_ddt_image(data):
    if len(data) > 3_000_000:
        raise ValueError("L'immagine deve essere inferiore a 3 MB.")
    with Image.open(io.BytesIO(data)) as img:
        img.verify()
    with Image.open(io.BytesIO(data)) as img:
        if img.format != "PNG" or img.width > 3000 or img.height > 3000:
            raise ValueError("Carica un PNG fino a 3000 × 3000 pixel.")
    return data


def signing_page(token):
    if not re.fullmatch(r"[0-9a-f]{64}", token or ""):
        st.error("Link di firma non valido.")
        return
    token_hash = hashlib.sha256(token.encode()).hexdigest()
    with db() as conn:
        request = conn.execute("SELECT * FROM document_sign_requests WHERE token_hash = ?", (token_hash,)).fetchone()
        doc = conn.execute("SELECT * FROM transport_documents WHERE id = ?", (request["document_id"],)).fetchone() if request else None
    if not request or not doc:
        st.error("Link di firma non valido.")
        return
    if request["signed_at"]:
        st.success("Documento già firmato. Grazie.")
        return
    if datetime.now() > datetime.fromisoformat(request["expires_at"]):
        st.error("Il link di firma è scaduto. Richiedi un nuovo link.")
        return
    st.title("Firma documento di trasporto")
    st.write(f"DDT {doc['number']} · {doc['document_date']}")
    st.write(f"Da {doc['departure'] or '—'} a {doc['destination'] or '—'}")
    st.write(f"Ruolo: {DDT_SIGNATURE_ROLES[request['role']]}")
    st.dataframe(pd.DataFrame(json.loads(doc["vehicles_json"] or "[]")), hide_index=True, use_container_width=True)
    st.download_button("Visualizza o scarica il DDT prima di firmare", ddt_pdf(doc),
                       f"DDT_{doc['number'].replace('/', '_')}.pdf", mime="application/pdf")
    from streamlit_drawable_canvas import st_canvas
    signer = st.text_input("Nome e cognome del firmatario *")
    st.caption("Traccia la firma con il dito, un pennino o il mouse. Usa il comando del riquadro per cancellare e rifare.")
    canvas = st_canvas(fill_color="rgba(255, 255, 255, 0)", stroke_width=3,
                       stroke_color="#12233c", background_color="#ffffff",
                       height=170, width=320, drawing_mode="freedraw", update_streamlit=True,
                       key=f"sign_{token_hash}")
    consent = st.checkbox("Confermo di aver letto il DDT e autorizzo l'apposizione della mia firma al documento.")
    if st.button("Firma il DDT", type="primary", disabled=not consent):
        if not signer.strip():
            st.error("Inserisci nome e cognome.")
        elif not canvas.json_data or not canvas.json_data.get("objects") or canvas.image_data is None:
            st.error("Traccia la firma prima di confermare.")
        else:
            image = Image.fromarray(canvas.image_data.astype("uint8"), "RGBA")
            buffer = io.BytesIO()
            image.save(buffer, format="PNG")
            now = datetime.now().isoformat(timespec="seconds")
            with db() as conn:
                updated = conn.execute("UPDATE document_sign_requests SET signed_at = ? WHERE token_hash = ? AND signed_at IS NULL AND expires_at > ?",
                                       (now, token_hash, now)).rowcount
                if updated:
                    conn.execute("""INSERT INTO document_assets(document_id,kind,image_data,signer_name,signed_at)
                                    VALUES(?,?,?,?,?) ON CONFLICT(document_id,kind) DO UPDATE SET
                                    image_data=excluded.image_data,signer_name=excluded.signer_name,signed_at=excluded.signed_at""",
                                 (doc["id"], request["role"], buffer.getvalue(), signer.strip(), now))
            if updated:
                st.success("Firma acquisita. Grazie.")
                st.rerun()
            else:
                st.error("Il link è già stato usato o è scaduto.")



def documents_page():
    st.header("Documenti · DDT")
    with db() as conn:
        docs = conn.execute("SELECT * FROM transport_documents ORDER BY document_date DESC, created_at DESC").fetchall()
    choices = {f"{r['number']} · {r['document_date']} · {r['destination'] or 'senza destinazione'}": r for r in docs}
    options = list(choices) + ["➕ Nuovo DDT"]
    selected_label = st.selectbox("Apri un DDT salvato o creane uno nuovo", options, key="ddt_document_selector_v2")
    selected = choices.get(selected_label)
    if selected:
        st.info(f"DDT {selected['number']} aperto. Modifica i campi qui sotto e premi «Salva modifiche». Più in basso trovi timbro, firme, PDF, Excel, stampa ed eliminazione.")
        quick_pdf = ddt_pdf(selected)
        left, right = st.columns(2)
        left.download_button("📄 Scarica PDF", quick_pdf,
                             f"DDT_{re.sub(r'[^A-Za-z0-9_-]', '_', selected['number'])}.pdf",
                             mime="application/pdf", use_container_width=True)
        right.download_button("📊 Scarica Excel", ddt_excel(selected),
                              f"DDT_{re.sub(r'[^A-Za-z0-9_-]', '_', selected['number'])}.xlsx",
                              mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", use_container_width=True)
    else:
        st.info("Compila il nuovo DDT e premi «Salva DDT». Dopo il salvataggio compariranno timbro, firme, PDF, Excel, stampa ed eliminazione.")
    if "ddt_selected_id" not in st.session_state or st.session_state.ddt_selected_id != (selected["id"] if selected else None):
        st.session_state.ddt_selected_id = selected["id"] if selected else None
        st.session_state.ddt_number = selected["number"] if selected else ""
        st.session_state.ddt_date = date.fromisoformat(selected["document_date"]) if selected else date.today()
    def value(key):
        return selected[key] or "" if selected else ""
    if not selected and not st.session_state.ddt_number:
        with db() as conn:
            st.session_state.ddt_number = next_ddt_number(conn, st.session_state.ddt_date.year)
    with st.form("ddt_form"):
        a,b = st.columns(2)
        number = a.text_input("Numero DDT * (automatico, modificabile)", key="ddt_number")
        document_date = b.date_input("Data documento *", key="ddt_date", format="DD/MM/YYYY")
        loading_date = st.date_input("Data carico", value=date.fromisoformat(value("loading_date")) if value("loading_date") else document_date, format="DD/MM/YYYY")
        a,b = st.columns(2)
        departure = a.text_input("Stazione di partenza", value=value("departure"))
        destination = b.text_input("Destinazione", value=value("destination"))
        a,b = st.columns(2)
        carrier = a.text_input("Vettore / Autisti", value=value("carrier"))
        truck_plate = b.text_input("Targa bisarca", value=value("truck_plate"))
        st.subheader("Veicoli trasportati")
        existing = json.loads(value("vehicles_json") or "[]")
        vehicle_rows = st.data_editor(pd.DataFrame(existing, columns=["marca", "modello", "targa", "note"]),
                                      num_rows="dynamic", hide_index=True, use_container_width=True,
                                      column_config={"marca":"Marca", "modello":"Modello", "targa":"Targa", "note":"Materiale a bordo / Note"})
        st.caption("Aggiungi o elimina righe con i comandi della tabella.")
        station_signature = st.text_input("Firma operatore stazione (nome)", value=value("station_signature"))
        driver_signature = st.text_input("Firma autista (nome)", value=value("driver_signature"))
        delivery_signature = st.text_input("Firma operatore stazione consegna / piazzale (nome)", value=value("delivery_signature"))
        saved = st.form_submit_button("Salva modifiche" if selected else "Salva DDT", type="primary")
    if saved:
        number = number.strip()
        if not number:
            st.error("Inserisci il numero del DDT.")
        else:
            vehicles = [{key: str(row.get(key) or "").strip() for key in ["marca","modello","targa","note"]}
                        for row in vehicle_rows.fillna("").to_dict("records")]
            vehicles = [row for row in vehicles if any(row.values())]
            now = datetime.now().isoformat(timespec="seconds")
            try:
                with db() as conn:
                    conn.execute("""INSERT INTO transport_documents
                        (id,number,document_date,loading_date,departure,destination,carrier,truck_plate,
                         station_signature,driver_signature,delivery_signature,vehicles_json,created_at,updated_at)
                        VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                        ON CONFLICT(id) DO UPDATE SET number=excluded.number,document_date=excluded.document_date,
                        loading_date=excluded.loading_date,departure=excluded.departure,destination=excluded.destination,
                        carrier=excluded.carrier,truck_plate=excluded.truck_plate,station_signature=excluded.station_signature,
                        driver_signature=excluded.driver_signature,delivery_signature=excluded.delivery_signature,
                        vehicles_json=excluded.vehicles_json,updated_at=excluded.updated_at""",
                        (selected["id"] if selected else uuid.uuid4().hex, number, document_date.isoformat(),
                         loading_date.isoformat(), departure.strip(), destination.strip(), carrier.strip(),
                         truck_plate.strip(), station_signature.strip(), driver_signature.strip(),
                         delivery_signature.strip(), json.dumps(vehicles, ensure_ascii=False), now, now))
                if selected:
                    changed_fields = {
                        "number": number, "document_date": document_date.isoformat(),
                        "loading_date": loading_date.isoformat(), "departure": departure.strip(),
                        "destination": destination.strip(), "carrier": carrier.strip(),
                        "truck_plate": truck_plate.strip(), "station_signature": station_signature.strip(),
                        "driver_signature": driver_signature.strip(), "delivery_signature": delivery_signature.strip(),
                    }
                    changed = any((selected[key] or "") != new_value for key, new_value in changed_fields.items())
                    changed = changed or json.loads(selected["vehicles_json"] or "[]") != vehicles
                    if changed:
                        with db() as conn:
                            conn.execute("DELETE FROM document_assets WHERE document_id = ? AND kind != 'stamp'", (selected["id"],))
                            conn.execute("DELETE FROM document_sign_requests WHERE document_id = ?", (selected["id"],))
                        st.warning("Il contenuto è cambiato: le firme precedenti e i link di firma sono stati annullati.")
                st.success(f"DDT {number} salvato.")
                st.rerun()
            except sqlite3.IntegrityError:
                st.error(f"Il numero {number} è già assegnato a un altro DDT.")
    if selected:
        with st.expander("Elimina questo DDT"):
            confirm = st.checkbox(f"Confermo l'eliminazione definitiva del DDT {selected['number']}", key=f"delete_ddt_{selected['id']}")
            if st.button("Elimina DDT e relative firme", disabled=not confirm, type="secondary"):
                with db() as conn:
                    conn.execute("DELETE FROM document_sign_requests WHERE document_id = ?", (selected["id"],))
                    conn.execute("DELETE FROM document_assets WHERE document_id = ?", (selected["id"],))
                    conn.execute("DELETE FROM transport_documents WHERE id = ?", (selected["id"],))
                st.session_state.ddt_selected_id = None
                st.session_state.ddt_sign_link = ""
                st.success("DDT eliminato.")
                st.rerun()
        st.subheader("Timbro e firme")
        assets = ddt_assets(selected["id"])
        kind = st.selectbox("Elemento da aggiungere", list(DDT_ASSET_KINDS),
                            format_func=lambda key: DDT_ASSET_KINDS[key])
        image_file = st.file_uploader("Carica timbro o firma in PNG", type=["png"], key=f"asset_{selected['id']}_{kind}")
        signer_name = st.text_input("Nome firmatario", key=f"signer_{selected['id']}_{kind}") if kind != "stamp" else ""
        if st.button("Salva immagine nel DDT"):
            if not image_file:
                st.error("Carica prima un'immagine PNG.")
            else:
                try:
                    image_bytes = validate_ddt_image(image_file.getvalue())
                    with db() as conn:
                        conn.execute("""INSERT INTO document_assets(document_id,kind,image_data,signer_name,signed_at)
                                        VALUES(?,?,?,?,?) ON CONFLICT(document_id,kind) DO UPDATE SET
                                        image_data=excluded.image_data,signer_name=excluded.signer_name,signed_at=excluded.signed_at""",
                                     (selected["id"], kind, image_bytes, signer_name.strip(),
                                      datetime.now().isoformat(timespec="seconds") if kind != "stamp" else None))
                    st.success("Immagine salvata nel DDT.")
                    st.rerun()
                except (ValueError, OSError) as exc:
                    st.error(str(exc))
        if kind in assets:
            st.image(assets[kind]["image_data"], width=200)
            if st.button(f"Rimuovi {DDT_ASSET_KINDS[kind].lower()}"):
                with db() as conn:
                    conn.execute("DELETE FROM document_assets WHERE document_id = ? AND kind = ?", (selected["id"],kind))
                st.rerun()
        st.subheader("Invio per la firma su telefono o tablet")
        with db() as conn:
            url_row = conn.execute("SELECT value FROM app_meta WHERE key = 'ddt_public_url'").fetchone()
        public_url = st.text_input("URL pubblico del gestionale (HTTPS)", value=url_row["value"] if url_row else "",
                                   help="Inserisci l'indirizzo HTTPS con cui il destinatario apre questa app.")
        st.caption("L'indirizzo viene salvato automaticamente quando generi il link.")
        role = st.selectbox("Chi deve firmare", list(DDT_SIGNATURE_ROLES), format_func=lambda key: DDT_SIGNATURE_ROLES[key])
        if st.button("Genera link di firma (valido 7 giorni)"):
            base_url = public_url.strip().rstrip("/")
            if not re.fullmatch(r"https://[^/\s?#]+(?:/[^?#]*)?", base_url):
                st.error("Inserisci l'URL HTTPS del gestionale nel campo qui sopra.")
            else:
                token = secrets.token_hex(32)
                with db() as conn:
                    conn.execute("INSERT INTO app_meta(key,value) VALUES('ddt_public_url',?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", (base_url,))
                    conn.execute("DELETE FROM document_sign_requests WHERE document_id = ? AND role = ? AND signed_at IS NULL", (selected["id"], role))
                    conn.execute("INSERT INTO document_sign_requests VALUES (?,?,?,?,NULL)",
                                 (hashlib.sha256(token.encode()).hexdigest(), selected["id"], role,
                                  (datetime.now()+timedelta(days=7)).isoformat(timespec="seconds")))
                st.session_state.ddt_sign_link = f"{base_url}?ddt_sign={token}"
                st.session_state.ddt_sign_document_id = selected["id"]
                st.success("Link pronto: puoi copiarlo o preparare l'email qui sotto.")
        if st.session_state.get("ddt_sign_link") and st.session_state.get("ddt_sign_document_id") == selected["id"]:
            st.code(st.session_state.ddt_sign_link, language=None)
            mailto = f"mailto:?subject={quote('Firma DDT '+selected['number'])}&body={quote('Apri il link per firmare il DDT: '+st.session_state.ddt_sign_link)}"
            st.link_button("Prepara email con il link", mailto)
            st.caption("Copia il link per inviarlo anche tramite messaggio. Generandone un altro, il precedente non sarà più utilizzabile.")
        st.subheader("Esporta e stampa")
        pdf_bytes = ddt_pdf(selected)
        st.download_button("Scarica DDT in Excel", ddt_excel(selected),
                           f"DDT_{re.sub(r'[^A-Za-z0-9_-]', '_', selected['number'])}.xlsx",
                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", use_container_width=True)
        st.download_button("Scarica DDT in PDF", pdf_bytes,
                           f"DDT_{re.sub(r'[^A-Za-z0-9_-]', '_', selected['number'])}.pdf",
                           mime="application/pdf", use_container_width=True)
        encoded = base64.b64encode(pdf_bytes).decode("ascii")
        st.components.v1.html(f"""<button onclick="let data=atob('{encoded}');let bytes=new Uint8Array(data.length);
            for(let i=0;i<data.length;i++)bytes[i]=data.charCodeAt(i);
            let url=URL.createObjectURL(new Blob([bytes],{{type:'application/pdf'}}));
            window.open(url,'_blank');" style="font-size:16px;padding:10px 16px;cursor:pointer">
            Apri PDF per stampare</button>""", height=55)
    if docs:
        st.subheader("Archivio DDT")
        st.dataframe(pd.DataFrame([{"Numero":r["number"],"Data":r["document_date"],
                                   "Partenza":r["departure"],"Destinazione":r["destination"],
                                   "Veicoli":len(json.loads(r["vehicles_json"] or "[]"))} for r in docs]),
                     hide_index=True, use_container_width=True)


def event_vehicles_to_pdf(frame, event_title):
    output = io.BytesIO()
    document = SimpleDocTemplate(
        output, pagesize=landscape(A4), rightMargin=8 * mm, leftMargin=8 * mm,
        topMargin=8 * mm, bottomMargin=8 * mm,
    )
    styles = getSampleStyleSheet()
    story = [Paragraph(f"Veicoli assegnati - {event_title}", styles["Title"]), Spacer(1, 4 * mm)]
    export = frame.copy()
    if "_id" in export.columns:
        export = export.drop(columns=["_id"])
    if "Stato" in export.columns:
        export["Stato"] = export["Data ritiro"].fillna("").astype(str).str.strip().map(lambda value: "RITIRATA" if value else "DA RITIRARE")
    export = export.fillna("").astype(str)
    data = [list(export.columns)] + export.values.tolist()
    widths = [23, 24, 23, 27, 30, 39, 25, 27, 39][:len(export.columns)]
    table = Table(data, repeatRows=1, colWidths=[width * mm for width in widths])
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCE6F1")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#172033")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F6F8FB")]),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
    ]
    if "Stato" in export.columns:
        status_column = list(export.columns).index("Stato")
        for row_number, status in enumerate(export["Stato"], start=1):
            style.append(("TEXTCOLOR", (status_column, row_number), (status_column, row_number), colors.green if status == "RITIRATA" else colors.red))
            style.append(("FONTNAME", (status_column, row_number), (status_column, row_number), "Helvetica-Bold"))
    table.setStyle(TableStyle(style))
    story.append(table)
    document.build(story)
    return output.getvalue()


def filters(frame):
    st.sidebar.markdown("### Filtri")
    years = sorted([int(x) for x in frame["start_date"].dt.year.dropna().unique()])
    selected_years = st.sidebar.multiselect("Anno", years, default=years)
    operators = sorted(x for x in frame["operator"].dropna().unique() if x)
    selected_operators = st.sidebar.multiselect("Operatore", operators)
    vehicles = sorted(x for x in frame["vehicle_type"].dropna().unique() if x)
    selected_vehicles = st.sidebar.multiselect("Veicolo", vehicles)
    sources = sorted(x for x in frame["source_base"].dropna().unique() if x)
    selected_sources = st.sidebar.multiselect("Fonte", sources)
    search = st.sidebar.text_input("Cerca RA, fonte o dettaglio")
    result = frame.copy()
    if selected_years:
        result = result[result["start_date"].dt.year.isin(selected_years)]
    if selected_operators:
        result = result[result["operator"].isin(selected_operators)]
    if selected_vehicles:
        result = result[result["vehicle_type"].isin(selected_vehicles)]
    if selected_sources:
        result = result[result["source_base"].isin(selected_sources)]
    if search:
        mask = result[["ra", "source_raw", "source_details", "operator"]].fillna("").apply(
            lambda col: col.str.contains(search, case=False, regex=False)
        ).any(axis=1)
        result = result[mask]
    return result


def metric_row(frame):
    rentals = len(frame)
    ancillary_count = int(frame["has_ancillary"].sum()) if rentals else 0
    revenue = float(frame["ancillary_cost"].fillna(0).sum()) if rentals else 0
    conversion = ancillary_count / rentals if rentals else 0
    average_ticket = revenue / ancillary_count if ancillary_count else 0
    average_rpd = float(frame.loc[frame["has_ancillary"], "rpd"].mean() or 0) if ancillary_count else 0
    cols = st.columns(6)
    cols[0].metric("Noleggi", f"{rentals:,}".replace(",", "."))
    cols[1].metric("Con ancillary", f"{ancillary_count:,}".replace(",", "."))
    cols[2].metric("Conversione", f"{conversion:.1%}")
    cols[3].metric("Ricavi ancillary", f"€ {revenue:,.2f}".replace(",", "X").replace(".", ",").replace("X", "."))
    cols[4].metric("Ticket medio", f"€ {average_ticket:.2f}".replace(".", ","))
    cols[5].metric("RPD medio", f"€ {average_rpd:.2f}".replace(".", ","))


def dashboard(frame):
    st.header("Dashboard ancillary")
    metric_row(frame)
    if frame.empty:
        st.info("Nessun dato corrisponde ai filtri selezionati.")
        return
    monthly = frame.dropna(subset=["start_date"]).copy()
    monthly["mese"] = monthly["start_date"].dt.to_period("M").astype(str)
    monthly = monthly.groupby("mese", as_index=False).agg(
        noleggi=("id", "count"), ricavi=("ancillary_cost", "sum"), conversione=("has_ancillary", "mean")
    )
    c1, c2 = st.columns(2)
    c1.plotly_chart(px.line(monthly, x="mese", y="ricavi", markers=True, title="Ricavi ancillary per mese"), use_container_width=True)
    by_operator = frame.groupby("operator", as_index=False).agg(
        noleggi=("id", "count"), ricavi=("ancillary_cost", "sum"), conversione=("has_ancillary", "mean")
    ).sort_values("ricavi", ascending=False)
    c2.plotly_chart(px.bar(by_operator, x="operator", y="ricavi", color="conversione", title="Ricavi e conversione per operatore"), use_container_width=True)
    c3, c4 = st.columns(2)
    by_source = frame.groupby("source_base", as_index=False).agg(
        noleggi=("id", "count"), ricavi=("ancillary_cost", "sum")
    ).sort_values("ricavi", ascending=False).head(15)
    c3.plotly_chart(px.bar(by_source, x="ricavi", y="source_base", orientation="h", title="Prime 15 fonti per ricavi"), use_container_width=True)
    by_vehicle = frame.groupby("vehicle_type", as_index=False).agg(
        noleggi=("id", "count"), ricavi=("ancillary_cost", "sum"), conversione=("has_ancillary", "mean")
    )
    c4.plotly_chart(px.bar(by_vehicle, x="vehicle_type", y=["noleggi", "ricavi"], barmode="group", title="Confronto CAR e VAN"), use_container_width=True)
    with st.expander("Controllo qualità dei dati"):
        duplicates = frame[frame.duplicated("ra", keep=False)].sort_values("ra")
        q1, q2, q3 = st.columns(3)
        q1.metric("RA duplicati", int(duplicates["ra"].nunique()))
        q2.metric("Fonte mancante", int(frame["source_raw"].fillna("").str.strip().eq("").sum()))
        q3.metric("Veicolo mancante", int(frame["vehicle_type"].fillna("").str.strip().eq("").sum()))
        if not duplicates.empty:
            st.dataframe(
                duplicates[["start_date", "ra", "operator", "source_raw"]],
                use_container_width=True,
                hide_index=True,
            )


def record_form(frame):
    st.header("Inserimento e modifica")
    records_by_id = {row.id: row for row in frame.itertuples()}
    options = [""] + list(records_by_id)
    selected = st.selectbox(
        "Record da modificare (lascia vuoto per inserirne uno nuovo)",
        options,
        format_func=lambda record_id: "" if not record_id else (
            f"{records_by_id[record_id].ra} | "
            f"{records_by_id[record_id].start_date.date() if pd.notna(records_by_id[record_id].start_date) else ''} | "
            f"{records_by_id[record_id].operator}"
        ),
    )
    current = None
    if selected:
        current = frame[frame["id"] == selected].iloc[0].to_dict()
    widget_prefix = (current or {}).get("id", "nuovo")
    rental_options = ["", "CORPORATE", "WALK-IN", "BROKER", "WEB/CNP", "REPLACEMENT", "MENSILE"]
    current_rental = upper((current or {}).get("rental_type", ""))
    rental_aliases = {
        "CORPARATE": "CORPORATE", "WEB - CNP ETC": "WEB/CNP", "WEB-CNP-ETC": "WEB/CNP",
        "ASSISTENZA": "REPLACEMENT", "FULL CREDIT, ASSISTENZA": "REPLACEMENT",
    }
    current_rental = rental_aliases.get(current_rental, current_rental)
    if current_rental not in rental_options:
        current_rental = ""

    a, b = st.columns(2)
    ra = a.text_input("RA (Rental Agreement) *", value=(current or {}).get("ra", ""), key=f"{widget_prefix}_ra")
    start_value = (current or {}).get("start_date")
    start_date = b.date_input(
        "Data inizio noleggio *",
        value=start_value.date() if pd.notna(start_value) else date.today(),
        key=f"{widget_prefix}_date",
    )

    d, e, f = st.columns(3)
    current_operator = upper((current or {}).get("operator", ""))
    operator_options = [""] + [row["name"] for row in configured_operators()]
    if current_operator and current_operator not in operator_options:
        operator_options.append(current_operator)
    operator = d.selectbox(
        "Operatore",
        operator_options,
        index=operator_options.index(current_operator) if current_operator in operator_options else 0,
        key=f"{widget_prefix}_operator",
    )
    vehicle_options = ["", "CAR", "VAN"]
    current_vehicle = upper((current or {}).get("vehicle_type", ""))
    if current_vehicle not in vehicle_options:
        current_vehicle = ""
    vehicle = e.selectbox(
        "Tipo veicolo", vehicle_options, index=vehicle_options.index(current_vehicle), key=f"{widget_prefix}_vehicle"
    )
    rental_type = f.selectbox(
        "Tipo noleggio *", rental_options, index=rental_options.index(current_rental), key=f"{widget_prefix}_rental"
    )

    ancillary_blocked = rental_type in {"CORPORATE", "MENSILE"}
    existing_days = int((current or {}).get("rental_days", 0) or 0)
    days = st.number_input(
        "Giorni noleggio *",
        min_value=0,
        step=1,
        value=0 if ancillary_blocked else max(existing_days, 1),
        disabled=ancillary_blocked,
        key=f"{widget_prefix}_days_{rental_type}",
    )

    car_ancillary = ["", "GOLD", "PLATINUM", "UP GOLD/PLATINUM", "GIÀ PRESENTE"]
    van_ancillary = ["", "NO PROBLEM", "SUPER VAN", "VAN PROTECTION", "CARGO VAN"]
    ancillary_options = car_ancillary if vehicle == "CAR" else van_ancillary if vehicle == "VAN" else [""]
    current_ancillary = upper((current or {}).get("ancillary", ""))
    if "PRESENTE" in current_ancillary and vehicle == "CAR":
        current_ancillary = "GIÀ PRESENTE"
    elif "UP" in current_ancillary and vehicle == "CAR":
        current_ancillary = "UP GOLD/PLATINUM"
    if current_ancillary not in ancillary_options:
        current_ancillary = ""

    g, h = st.columns(2)
    ancillary = g.selectbox(
        "Tipo ancillary",
        ancillary_options,
        index=ancillary_options.index(current_ancillary),
        disabled=ancillary_blocked or not vehicle,
        key=f"{widget_prefix}_ancillary_{vehicle}_{rental_type}",
    )
    existing_cost = (current or {}).get("ancillary_cost")
    existing_daily_cost = (
        float(existing_cost) / existing_days
        if pd.notna(existing_cost) and existing_days > 0
        else 0.0
    )
    daily_cost = h.number_input(
        "Costo giornaliero ancillary (IVA esclusa)",
        min_value=0.0,
        step=0.01,
        value=0.0 if ancillary_blocked else round(existing_daily_cost, 2),
        disabled=ancillary_blocked or not ancillary,
        key=f"{widget_prefix}_cost_{vehicle}_{rental_type}_{ancillary}",
    )
    calculated_total = round(daily_cost * days, 2) if not ancillary_blocked and days else 0.0
    if ancillary_blocked:
        st.info(f"Per i noleggi {rental_type} non vengono inseriti giorni, ancillary o costi ancillary.")
    else:
        st.info(f"Totale ancillary calcolato: € {daily_cost:.2f} × {days} giorni = € {calculated_total:.2f}")
    notes = st.text_area("Note", value=(current or {}).get("notes", ""), key=f"{widget_prefix}_notes")
    submitted = st.button("Salva record", type="primary", use_container_width=True)
    if submitted:
        if not ra or not rental_type:
            st.error("Inserisci almeno RA e tipo noleggio.")
        elif start_date < ANCILLARY_START_DATE:
            st.error("Lo storico ancillary parte dal 01/10/2026. Inserisci una data uguale o successiva.")
        elif not ancillary_blocked and days <= 0:
            st.error("Inserisci i giorni di noleggio.")
        else:
            save_record({
                "id": (current or {}).get("id"), "created_at": (current or {}).get("created_at"),
                "ra": ra, "start_date": start_date.isoformat(), "rental_days": days,
                "source_raw": (current or {}).get("source_raw", ""),
                "operator": operator, "vehicle_type": vehicle,
                "ancillary": ancillary,
                "ancillary_cost": calculated_total if calculated_total > 0 else None,
                "rental_type": rental_type, "notes": notes,
            })
            st.success("Record salvato e normalizzato.")
            st.rerun()
    if current:
        st.divider()
        confirm = st.checkbox(f"Confermo l’eliminazione definitiva di {current['ra']}")
        if st.button("Elimina record", disabled=not confirm, type="secondary"):
            with db() as conn:
                conn.execute("DELETE FROM rentals WHERE id = ?", (current["id"],))
            st.success("Record eliminato.")
            st.rerun()


def archive(frame):
    st.header("Archivio noleggi")
    display = frame[[
        "start_date", "ra", "rental_days", "source_base", "source_details",
        "operator", "vehicle_type", "ancillary", "ancillary_cost", "rpd", "rental_type",
    ]].copy()
    display.columns = [
        "Data", "RA", "Giorni", "Fonte", "Dettaglio fonte", "Operatore",
        "Veicolo", "Ancillary", "Costo", "RPD", "Tipo noleggio",
    ]
    st.dataframe(
        display,
        use_container_width=True,
        hide_index=True,
        column_config={
            "Data": st.column_config.DateColumn(format="DD/MM/YYYY"),
            "Costo": st.column_config.NumberColumn(format="€ %.2f"),
            "RPD": st.column_config.NumberColumn(format="€ %.2f"),
        },
    )
    st.download_button("Esporta archivio filtrato in Excel", to_excel(frame), "archivio_ancillary.xlsx", use_container_width=True)


def contract_filters(frame):
    if frame.empty:
        return frame
    c1, c2, c3, c4 = st.columns(4)
    periods = sorted(x for x in frame["import_period"].dropna().unique() if x)
    selected_periods = c1.multiselect("Periodo RA", periods, default=periods)
    operators = sorted(x for x in frame["operator"].dropna().unique() if x)
    selected_operators = c2.multiselect("Operatore RA", operators)
    sources = sorted(x for x in frame["source_base"].dropna().unique() if x)
    selected_sources = c3.multiselect("Fonte RA", sources)
    groups = sorted(x for x in frame["assigned_group"].dropna().unique() if x)
    selected_groups = c4.multiselect("Gruppo assegnato", groups)
    c5, c6, c7 = st.columns(3)
    channels = sorted(x for x in frame["rental_channel"].dropna().unique() if x)
    selected_channels = c5.multiselect("Tipo noleggio", channels)
    terms = [x for x in ["GIORNALIERO", "MENSILE"] if x in set(frame["rental_term"])]
    selected_terms = c6.multiselect("Durata noleggio", terms)
    vehicles = [x for x in ["CAR", "VAN"] if x in set(frame["contract_vehicle"])]
    selected_vehicles = c7.multiselect("Tipo veicolo", vehicles)
    result = frame.copy()
    if selected_periods:
        result = result[result["import_period"].isin(selected_periods)]
    if selected_operators:
        result = result[result["operator"].isin(selected_operators)]
    if selected_sources:
        result = result[result["source_base"].isin(selected_sources)]
    if selected_groups:
        result = result[result["assigned_group"].isin(selected_groups)]
    if selected_channels:
        result = result[result["rental_channel"].isin(selected_channels)]
    if selected_terms:
        result = result[result["rental_term"].isin(selected_terms)]
    if selected_vehicles:
        result = result[result["contract_vehicle"].isin(selected_vehicles)]
    return result


def contracts_dashboard(frame):
    st.header("Analisi contratti RA")
    frame = contract_filters(frame)
    if frame.empty:
        st.info("Nessun contratto disponibile per i filtri selezionati.")
        return
    contracts = len(frame)
    days = int(frame["duration_days"].fillna(0).sum())
    value = float(frame["contract_value"].fillna(0).sum())
    ancillary = float(frame["ancillary_value"].fillna(0).sum())
    rpd = value / days if days else 0
    ancillary_rpd = ancillary / days if days else 0
    penetration = float(frame["has_ancillary_ra"].mean()) if contracts else 0
    cols = st.columns(7)
    cols[0].metric("Contratti", f"{contracts:,}".replace(",", "."))
    cols[1].metric("Giorni", f"{days:,}".replace(",", "."))
    cols[2].metric("Valore contratti", f"€ {value:,.0f}".replace(",", "."))
    cols[3].metric("RPD contratti", f"€ {rpd:.2f}".replace(".", ","))
    cols[4].metric("Valore ancillary", f"€ {ancillary:,.0f}".replace(",", "."))
    cols[5].metric("RPD ancillary", f"€ {ancillary_rpd:.2f}".replace(".", ","))
    cols[6].metric("Contratti con ancillary", f"{penetration:.1%}")

    daily = frame.copy()
    daily["data"] = daily["contract_date"].dt.date
    by_day = daily.groupby("data", as_index=False).agg(
        contratti=("id", "count"), valore=("contract_value", "sum"), ancillary=("ancillary_value", "sum")
    )
    c1, c2 = st.columns(2)
    c1.plotly_chart(px.line(by_day, x="data", y=["valore", "ancillary"], markers=True, title="Valore contratti e ancillary per giorno"), use_container_width=True)
    by_operator = frame.groupby("operator", as_index=False).agg(
        contratti=("id", "count"), giorni=("duration_days", "sum"),
        valore=("contract_value", "sum"), ancillary=("ancillary_value", "sum")
    )
    by_operator["rpd"] = by_operator["valore"].div(by_operator["giorni"].replace(0, pd.NA))
    by_operator["ancillary_rpd"] = by_operator["ancillary"].div(by_operator["giorni"].replace(0, pd.NA))
    c2.plotly_chart(px.bar(by_operator.sort_values("valore", ascending=False), x="operator", y="valore", color="ancillary_rpd", title="Valore e RPD ancillary per operatore"), use_container_width=True)

    c3, c4 = st.columns(2)
    by_source = frame.groupby("source_base", as_index=False).agg(
        contratti=("id", "count"), valore=("contract_value", "sum"), ancillary=("ancillary_value", "sum")
    ).sort_values("valore", ascending=False).head(15)
    c3.plotly_chart(px.bar(by_source, x="valore", y="source_base", color="ancillary", orientation="h", title="Prime fonti per valore contratto"), use_container_width=True)
    groups = frame.groupby(["requested_group", "assigned_group"], as_index=False).size().sort_values("size", ascending=False).head(20)
    groups["passaggio"] = groups["requested_group"] + " → " + groups["assigned_group"]
    c4.plotly_chart(px.bar(groups, x="size", y="passaggio", orientation="h", title="Gruppo richiesto e assegnato"), use_container_width=True)

    breakdown = frame.groupby(
        ["rental_channel", "rental_term", "contract_vehicle"], as_index=False
    ).agg(
        contratti=("id", "count"), giorni=("duration_days", "sum"),
        valore=("contract_value", "sum"), ancillary=("ancillary_value", "sum"),
    )
    breakdown["rpd"] = breakdown["valore"].div(breakdown["giorni"].replace(0, pd.NA))
    breakdown["ancillary_rpd"] = breakdown["ancillary"].div(breakdown["giorni"].replace(0, pd.NA))
    breakdown["combinazione"] = breakdown["rental_channel"] + " - " + breakdown["rental_term"]
    c5, c6 = st.columns(2)
    c5.plotly_chart(
        px.bar(
            breakdown, x="combinazione", y="contratti", color="contract_vehicle",
            barmode="group", title="Contratti per tipo, durata e veicolo",
        ),
        use_container_width=True,
    )
    vehicle_summary = frame.groupby("contract_vehicle", as_index=False).agg(
        contratti=("id", "count"), valore=("contract_value", "sum"), ancillary=("ancillary_value", "sum")
    )
    c6.plotly_chart(
        px.bar(
            vehicle_summary, x="contract_vehicle", y=["valore", "ancillary"],
            barmode="group", title="Confronto CAR e VAN",
        ),
        use_container_width=True,
    )

    st.subheader("Risultati per tipo di noleggio, durata e veicolo")
    st.dataframe(breakdown, use_container_width=True, hide_index=True)

    st.subheader("Risultati per operatore")
    st.dataframe(by_operator.sort_values("valore", ascending=False), use_container_width=True, hide_index=True)


def ancillary_ra_dashboard(frame):
    st.header("Analisi ancillary RA")
    st.caption("Tutti i valori provengono dai report Analisi contratti e seguono i filtri selezionati, incluso il gruppo assegnato.")
    frame = contract_filters(frame)
    if frame.empty:
        st.info("Nessun dato ancillary disponibile per i filtri selezionati.")
        return

    contracts = len(frame)
    with_ancillary = int(frame["has_ancillary_ra"].sum())
    days = int(frame["duration_days"].fillna(0).sum())
    ancillary_value = float(frame["ancillary_value"].fillna(0).sum())
    ancillary_rpd = ancillary_value / days if days else 0
    penetration = with_ancillary / contracts if contracts else 0
    ticket = ancillary_value / with_ancillary if with_ancillary else 0
    cols = st.columns(6)
    cols[0].metric("Contratti analizzati", f"{contracts:,}".replace(",", "."))
    cols[1].metric("Con ancillary", f"{with_ancillary:,}".replace(",", "."))
    cols[2].metric("Penetrazione", f"{penetration:.1%}")
    cols[3].metric("Valore ancillary", f"€ {ancillary_value:,.0f}".replace(",", "."))
    cols[4].metric("RPD ancillary", f"€ {ancillary_rpd:.2f}".replace(".", ","))
    cols[5].metric("Ticket medio ancillary", f"€ {ticket:.2f}".replace(".", ","))

    daily = frame.copy()
    daily["data"] = daily["contract_date"].dt.date
    by_day = daily.groupby("data", as_index=False).agg(
        valore_ancillary=("ancillary_value", "sum"),
        contratti=("id", "count"),
        contratti_con_ancillary=("has_ancillary_ra", "sum"),
    )
    by_day["penetrazione"] = by_day["contratti_con_ancillary"].div(by_day["contratti"].replace(0, pd.NA))

    c1, c2 = st.columns(2)
    c1.plotly_chart(
        px.line(by_day, x="data", y="valore_ancillary", markers=True, title="Valore ancillary per giorno"),
        use_container_width=True,
    )
    by_operator = frame.groupby("operator", as_index=False).agg(
        contratti=("id", "count"), giorni=("duration_days", "sum"),
        valore_ancillary=("ancillary_value", "sum"), contratti_con_ancillary=("has_ancillary_ra", "sum"),
    )
    by_operator["rpd_ancillary"] = by_operator["valore_ancillary"].div(by_operator["giorni"].replace(0, pd.NA))
    by_operator["penetrazione"] = by_operator["contratti_con_ancillary"].div(by_operator["contratti"].replace(0, pd.NA))
    c2.plotly_chart(
        px.bar(by_operator, x="operator", y="valore_ancillary", color="rpd_ancillary", title="Valore e RPD ancillary per operatore"),
        use_container_width=True,
    )

    c3, c4 = st.columns(2)
    by_group = frame.groupby("assigned_group", as_index=False).agg(
        contratti=("id", "count"), giorni=("duration_days", "sum"),
        valore_ancillary=("ancillary_value", "sum"), contratti_con_ancillary=("has_ancillary_ra", "sum"),
    )
    by_group["rpd_ancillary"] = by_group["valore_ancillary"].div(by_group["giorni"].replace(0, pd.NA))
    by_group["penetrazione"] = by_group["contratti_con_ancillary"].div(by_group["contratti"].replace(0, pd.NA))
    c3.plotly_chart(
        px.bar(by_group.sort_values("valore_ancillary", ascending=False), x="assigned_group", y="valore_ancillary", color="penetrazione", title="Ancillary per gruppo assegnato"),
        use_container_width=True,
    )
    by_mix = frame.groupby(["rental_channel", "rental_term", "contract_vehicle"], as_index=False).agg(
        contratti=("id", "count"), valore_ancillary=("ancillary_value", "sum"),
    )
    by_mix["combinazione"] = by_mix["rental_channel"] + " - " + by_mix["rental_term"]
    c4.plotly_chart(
        px.bar(by_mix, x="combinazione", y="valore_ancillary", color="contract_vehicle", barmode="group", title="Ancillary per tipo, durata e veicolo"),
        use_container_width=True,
    )

    st.subheader("Statistiche ancillary per gruppo assegnato")
    st.dataframe(by_group.sort_values("valore_ancillary", ascending=False), use_container_width=True, hide_index=True)
    st.subheader("Statistiche ancillary per operatore")
    st.dataframe(by_operator.sort_values("valore_ancillary", ascending=False), use_container_width=True, hide_index=True)


def contracts_archive(frame):
    st.header("Archivio contratti RA")
    frame = contract_filters(frame)
    display = frame[[
        "contract_date", "ra", "operator", "client", "duration_days",
        "rental_channel", "rental_term", "contract_vehicle",
        "requested_group", "assigned_group", "source_base", "source_details",
        "contract_value", "kpi1_rpd", "ancillary_value", "ancillary_rpd", "import_file",
    ]].copy()
    display.columns = [
        "Data", "RA", "Operatore", "Cliente", "Giorni",
        "Tipo noleggio", "Durata noleggio", "Veicolo", "Gruppo richiesto",
        "Gruppo assegnato", "Fonte", "Dettaglio fonte", "Valore contratto",
        "RPD contratto", "Valore ancillary", "RPD ancillary", "File importato",
    ]
    st.dataframe(display, use_container_width=True, hide_index=True)


def combined_analysis(contracts, ancillary):
    st.header("Analisi incrociata RA e ancillary")
    if contracts.empty or ancillary.empty:
        st.info("Servono sia contratti RA sia dati ancillary per eseguire il confronto.")
        return
    ancillary = ancillary.copy()
    contracts = contracts.copy()
    ancillary["analysis_year"] = ancillary["start_date"].dt.year
    contracts["analysis_year"] = contracts["contract_date"].dt.year.fillna(contracts["start_date"].dt.year)
    ancillary_summary = ancillary.groupby(["ra", "analysis_year"], as_index=False).agg(
        ancillary_modulo=("ancillary_cost", "sum"), righe_modulo=("id", "count")
    )
    joined = contracts.merge(ancillary_summary, on=["ra", "analysis_year"], how="left")
    joined["ancillary_modulo"] = joined["ancillary_modulo"].fillna(0)
    joined["differenza"] = joined["ancillary_value"].fillna(0) - joined["ancillary_modulo"]
    matched = joined["righe_modulo"].notna().sum()
    c1, c2, c3, c4 = st.columns(4)
    c1.metric("RA nel periodo", len(joined))
    c2.metric("RA collegati al modulo", int(matched))
    c3.metric("Copertura collegamento", f"{matched / len(joined):.1%}" if len(joined) else "0%")
    c4.metric("Differenza complessiva", f"€ {joined['differenza'].sum():.2f}".replace(".", ","))
    st.plotly_chart(
        px.scatter(joined, x="ancillary_modulo", y="ancillary_value", color="operator", hover_data=["ra", "source_base"], title="Confronto ancillary modulo e report RA"),
        use_container_width=True,
    )
    anomalies = joined[joined["differenza"].abs() > 0.02][[
        "ra", "operator", "source_base", "ancillary_modulo", "ancillary_value", "differenza"
    ]].sort_values("differenza", key=lambda x: x.abs(), ascending=False)
    st.subheader("Differenze da controllare")
    st.dataframe(anomalies, use_container_width=True, hide_index=True)


def damage_filters(frame):
    if frame.empty:
        return frame
    c1, c2, c3, c4 = st.columns(4)
    years = sorted(frame["control_date"].dropna().dt.year.unique().astype(int).tolist())
    selected_years = c1.multiselect("Anno controllo", years, default=years)
    vehicles = sorted(x for x in frame["vehicle_category"].dropna().unique() if x)
    selected_vehicles = c2.multiselect("Categoria veicolo", vehicles)
    modes = sorted(x for x in frame["charge_mode"].dropna().unique() if x)
    selected_modes = c3.multiselect("Modalità addebito", modes)
    operators = sorted(x for x in frame["operator"].dropna().unique() if x)
    selected_operators = c4.multiselect("Operatore", operators)
    c5, c6, c7 = st.columns(3)
    categories = sorted(x for x in frame["damage_category"].dropna().unique() if x)
    selected_categories = c5.multiselect("Tipo danno", categories)
    statuses = sorted(x for x in frame["payment_status"].dropna().unique() if x)
    selected_statuses = c6.multiselect("Stato addebito", statuses)
    search = c7.text_input("Cerca RA o descrizione")
    result = frame.copy()
    if selected_years:
        result = result[result["control_date"].dt.year.isin(selected_years)]
    if selected_vehicles:
        result = result[result["vehicle_category"].isin(selected_vehicles)]
    if selected_modes:
        result = result[result["charge_mode"].isin(selected_modes)]
    if selected_operators:
        result = result[result["operator"].isin(selected_operators)]
    if selected_categories:
        result = result[result["damage_category"].isin(selected_categories)]
    if selected_statuses:
        result = result[result["payment_status"].isin(selected_statuses)]
    if search:
        needle = upper(search)
        compact_needle = re.sub(r"[^A-Z0-9]", "", needle)
        searchable = result[["ra", "description", "operator"]].fillna("")
        normal_match = searchable.apply(
            lambda column: column.str.upper().str.contains(needle, regex=False)
        ).any(axis=1)
        compact_match = searchable.apply(
            lambda column: column.map(
                lambda value: compact_needle in re.sub(r"[^A-Z0-9]", "", upper(value))
            )
        ).any(axis=1)
        mask = normal_match | compact_match
        result = result[mask]
    return result


def damage_dashboard(frame):
    st.header("Analisi addebito danni")
    frame = damage_filters(frame)
    if frame.empty:
        st.info("Nessuna segnalazione danni disponibile per i filtri selezionati.")
        return
    total = len(frame)
    vans = int(frame["vehicle_category"].eq("VAN").sum())
    cars = int(frame["vehicle_category"].eq("CAR").sum())
    photo_coverage = float(frame["has_photo"].mean()) if total else 0
    charged = float(frame["amount"].fillna(0).sum())
    collected = float(frame.loc[frame["payment_status"].eq("INCASSATO"), "amount"].fillna(0).sum())
    cols = st.columns(6)
    cols[0].metric("Segnalazioni", total)
    cols[1].metric("CAR", cars)
    cols[2].metric("VAN", vans)
    cols[3].metric("Con foto", f"{photo_coverage:.1%}")
    cols[4].metric("Importo addebitato", f"€ {charged:,.2f}".replace(",", "X").replace(".", ",").replace("X", "."))
    cols[5].metric("Importo incassato", f"€ {collected:,.2f}".replace(",", "X").replace(".", ",").replace("X", "."))

    monthly = frame.dropna(subset=["control_date"]).copy()
    monthly["mese"] = monthly["control_date"].dt.to_period("M").astype(str)
    monthly = monthly.groupby("mese", as_index=False).agg(segnalazioni=("id", "count"), importo=("amount", "sum"))
    c1, c2 = st.columns(2)
    c1.plotly_chart(px.line(monthly, x="mese", y="segnalazioni", markers=True, title="Segnalazioni per mese"), use_container_width=True)
    by_operator = frame.groupby("operator", as_index=False).agg(segnalazioni=("id", "count"), importo=("amount", "sum"))
    c2.plotly_chart(px.bar(by_operator, x="operator", y="segnalazioni", color="importo", title="Danni per operatore"), use_container_width=True)
    c3, c4 = st.columns(2)
    by_mode = frame.groupby("charge_mode", as_index=False).size().rename(columns={"size": "segnalazioni"})
    c3.plotly_chart(px.bar(by_mode, x="charge_mode", y="segnalazioni", title="Danni per modalità di addebito"), use_container_width=True)
    by_category = frame.groupby(["damage_category", "vehicle_category"], as_index=False).size().rename(columns={"size": "segnalazioni"})
    c4.plotly_chart(px.bar(by_category, x="damage_category", y="segnalazioni", color="vehicle_category", barmode="group", title="Tipologia danni CAR e VAN"), use_container_width=True)
    st.subheader("Riepilogo per operatore")
    st.dataframe(by_operator.sort_values("segnalazioni", ascending=False), use_container_width=True, hide_index=True)


def damage_archive(frame):
    st.header("Archivio addebito danni")
    frame = damage_filters(frame)
    if frame.empty:
        st.info("Nessuna segnalazione presente.")
        return
    display = frame[[
        "control_date", "ra", "vehicle_category", "charge_mode", "operator",
        "damage_category", "description", "photo_count", "amount", "payment_status", "photo_url", "notes",
    ]].copy()
    display.columns = [
        "Data controllo", "RA", "Veicolo", "Modalità", "Operatore", "Tipo danno",
        "Descrizione", "Numero foto", "Importo", "Stato", "Collegamento foto", "Note",
    ]
    st.dataframe(
        display, use_container_width=True, hide_index=True,
        column_config={
            "Data controllo": st.column_config.DateColumn(format="DD/MM/YYYY"),
            "Importo": st.column_config.NumberColumn(format="€ %.2f"),
            "Collegamento foto": st.column_config.LinkColumn(display_text="Apri foto"),
        },
    )
    st.download_button(
        "Esporta archivio danni in Excel", table_to_excel(display, "Addebito danni"),
        "archivio_addebito_danni.xlsx", use_container_width=True,
    )

    records_with_photos = frame[frame["photo_count"].gt(0)]
    if not records_with_photos.empty:
        st.subheader("Fotografie delle segnalazioni")
        photo_record_id = st.selectbox(
            "Seleziona RA",
            records_with_photos["id"].tolist(),
            format_func=lambda item_id: (
                f"{records_with_photos.loc[records_with_photos['id'].eq(item_id), 'ra'].iloc[0]} — "
                f"{int(records_with_photos.loc[records_with_photos['id'].eq(item_id), 'photo_count'].iloc[0])} foto"
            ),
        )
        photos = damage_photos(photo_record_id)
        columns = st.columns(min(3, len(photos)))
        for index, photo in enumerate(photos):
            with columns[index % len(columns)]:
                st.image(photo["file_data"], caption=photo["file_name"], use_container_width=True)
                st.download_button(
                    "Scarica", photo["file_data"], photo["file_name"],
                    mime=photo["mime_type"] or "image/jpeg", key=f"download_photo_{photo['id']}",
                    use_container_width=True,
                )


def quick_damage_checkin():
    st.header("Check-in rapido danni")
    st.caption("Compila i dati essenziali e allega più fotografie direttamente da telefono o tablet.")

    linked_operator = current_operator() if not user_is_admin() else ""
    operator_options = [row["name"] for row in configured_operators()]
    if linked_operator and linked_operator not in operator_options:
        operator_options.append(linked_operator)

    with st.form("quick_damage_checkin", clear_on_submit=True):
        ra = st.text_input("RA *", placeholder="Esempio: TOR-12345")
        c1, c2 = st.columns(2)
        vehicle = c1.selectbox("Categoria veicolo *", ["", "CAR", "VAN"])
        charge_mode = c2.selectbox(
            "Modalità di addebito *",
            ["IN PRESENZA", "KEY BOX", "CLIENTE NON ASPETTA (AFTER RENTAL)", "ALTRO"],
        )
        if linked_operator:
            operator = st.selectbox(
                "Operatore responsabile *", operator_options,
                index=operator_options.index(linked_operator), disabled=True,
            )
            st.caption("Operatore compilato automaticamente dall’account personale.")
        else:
            operator = st.selectbox("Operatore responsabile *", [""] + operator_options)
        description = st.text_area(
            "Descrizione sintetica del danno *", placeholder="Indica posizione e tipo di danno", height=100,
        )
        photos = st.file_uploader(
            "Fotografie del danno",
            type=["jpg", "jpeg", "png", "webp", "heic"],
            accept_multiple_files=True,
            help="Puoi scattare o selezionare più fotografie dal telefono.",
        )
        notes = st.text_area("Note facoltative", height=70)
        submitted = st.form_submit_button("Salva check-in", type="primary", use_container_width=True)

    if submitted:
        if not all([ra, vehicle, charge_mode, operator, description]):
            st.error("Compila RA, categoria veicolo, modalità, operatore e descrizione.")
            return
        if not photos:
            st.error("Allega almeno una fotografia del danno prima di salvare.")
            return
        if len(photos or []) > 10:
            st.error("Puoi allegare al massimo 10 fotografie per check-in.")
            return
        oversized = [photo.name for photo in photos or [] if len(photo.getvalue()) > 8 * 1024 * 1024]
        if oversized:
            st.error("Ogni fotografia deve essere inferiore a 8 MB: " + ", ".join(oversized))
            return
        damage_id = save_damage({
            "control_date": date.today().isoformat(), "ra": ra,
            "vehicle_category": vehicle, "charge_mode": charge_mode,
            "operator": operator, "description": description, "photo_url": "",
            "amount": None, "payment_status": "DA VERIFICARE", "notes": notes,
        })
        photo_total = save_damage_photos(damage_id, photos)
        st.success(f"Check-in {normalize_ra(ra)} salvato con {photo_total} fotografie.")
        st.rerun()


def damage_form(frame):
    st.header("Inserimento e modifica addebito danni")
    records = {row.id: row for row in frame.itertuples()}
    selected = st.selectbox(
        "Segnalazione da modificare (lascia vuoto per inserirne una nuova)",
        [""] + list(records),
        format_func=lambda record_id: "" if not record_id else (
            f"{records[record_id].ra} | "
            f"{records[record_id].control_date.date() if pd.notna(records[record_id].control_date) else ''} | "
            f"{records[record_id].operator}"
        ),
    )
    current = frame[frame["id"] == selected].iloc[0].to_dict() if selected else {}
    prefix = current.get("id", "nuovo_danno")
    c1, c2, c3 = st.columns(3)
    ra = c1.text_input("RA *", value=current.get("ra", ""), key=f"{prefix}_damage_ra")
    current_date = current.get("control_date")
    control_date = c2.date_input(
        "Data del controllo *",
        value=current_date.date() if pd.notna(current_date) else date.today(),
        key=f"{prefix}_damage_date",
    )
    vehicle_options = ["", "CAR", "VAN"]
    current_vehicle = upper(current.get("vehicle_category", ""))
    vehicle = c3.selectbox(
        "Categoria veicolo *", vehicle_options,
        index=vehicle_options.index(current_vehicle) if current_vehicle in vehicle_options else 0,
        key=f"{prefix}_damage_vehicle",
    )
    c4, c5 = st.columns(2)
    mode_options = ["", "IN PRESENZA", "KEY BOX", "CLIENTE NON ASPETTA (AFTER RENTAL)", "ALTRO"]
    current_mode = upper(current.get("charge_mode", ""))
    if current_mode and current_mode not in mode_options:
        mode_options.append(current_mode)
    charge_mode = c4.selectbox(
        "Modalità di addebito *", mode_options,
        index=mode_options.index(current_mode) if current_mode in mode_options else 0,
        key=f"{prefix}_damage_mode",
    )
    operator_options = [""] + [row["name"] for row in configured_operators()]
    current_operator_value = upper(current.get("operator", ""))
    linked_operator = current_operator() if not user_is_admin() else ""
    if linked_operator:
        current_operator_value = linked_operator
    if current_operator_value and current_operator_value not in operator_options:
        operator_options.append(current_operator_value)
    operator = c5.selectbox(
        "Operatore responsabile *", operator_options,
        index=operator_options.index(current_operator_value) if current_operator_value in operator_options else 0,
        key=f"{prefix}_damage_operator",
        disabled=bool(linked_operator),
    )
    if linked_operator:
        c5.caption("Compilato automaticamente dall’utente collegato.")
    description = st.text_area("Descrizione sintetica del danno *", value=current.get("description", ""), key=f"{prefix}_damage_description")
    photo_url = st.text_input("Collegamento foto danno", value=current.get("photo_url", ""), key=f"{prefix}_damage_photo")
    existing_photo_total = len(damage_photos(current["id"])) if current else 0
    new_photos = st.file_uploader(
        "Aggiungi fotografie",
        type=["jpg", "jpeg", "png", "webp", "heic"], accept_multiple_files=True,
        key=f"{prefix}_damage_new_photos",
        help=f"Fotografie già archiviate: {existing_photo_total}. Le nuove immagini verranno aggiunte.",
    )
    c6, c7 = st.columns(2)
    amount_value = current.get("amount")
    amount = c6.number_input(
        "Importo addebitato", min_value=0.0, step=10.0,
        value=0.0 if pd.isna(amount_value) else float(amount_value), key=f"{prefix}_damage_amount",
    )
    status_options = ["DA VERIFICARE", "ADDEBITATO", "INCASSATO", "ANNULLATO"]
    current_status = upper(current.get("payment_status", "")) or "DA VERIFICARE"
    status = c7.selectbox(
        "Stato addebito", status_options,
        index=status_options.index(current_status) if current_status in status_options else 0,
        key=f"{prefix}_damage_status",
    )
    notes = st.text_area("Note", value=current.get("notes", ""), key=f"{prefix}_damage_notes")
    if st.button("Salva segnalazione danno", type="primary", use_container_width=True):
        if not all([ra, vehicle, charge_mode, operator, description]):
            st.error("Compila RA, categoria veicolo, modalità, operatore e descrizione.")
        else:
            saved_damage_id = save_damage({
                "id": current.get("id"), "source_key": current.get("source_key"),
                "submitted_at": current.get("submitted_at"), "control_date": control_date.isoformat(),
                "ra": ra, "vehicle_category": vehicle, "charge_mode": charge_mode,
                "operator": operator, "description": description, "photo_url": photo_url,
                "amount": amount if amount > 0 else None, "payment_status": status, "notes": notes,
            })
            added_photos = save_damage_photos(saved_damage_id, new_photos)
            st.success(f"Segnalazione danno salvata. Nuove fotografie aggiunte: {added_photos}.")
            st.rerun()
    if current:
        confirm = st.checkbox(f"Confermo l’eliminazione definitiva della segnalazione {current['ra']}")
        if st.button("Elimina segnalazione danno", disabled=not confirm, use_container_width=True):
            with db() as conn:
                conn.execute("DELETE FROM damage_photos WHERE damage_id = ?", (current["id"],))
                conn.execute("DELETE FROM damage_charges WHERE id = ?", (current["id"],))
            st.success("Segnalazione eliminata.")
            st.rerun()


def cash_filters(frame):
    if frame.empty:
        return frame
    st.sidebar.subheader("Filtri cassa")
    years = sorted(frame["movement_date"].dropna().dt.year.unique().tolist(), reverse=True)
    selected_years = st.sidebar.multiselect("Anno cassa", years, default=years[:1])
    movement_types = sorted(frame["movement_type"].dropna().unique().tolist())
    selected_types = st.sidebar.multiselect("Tipo movimento", movement_types)
    methods = sorted(frame["payment_method"].dropna().unique().tolist())
    selected_methods = st.sidebar.multiselect("Metodo", methods)
    search = st.sidebar.text_input("Cerca RA o note")
    result = frame.copy()
    if selected_years:
        result = result[result["movement_date"].dt.year.isin(selected_years)]
    if selected_types:
        result = result[result["movement_type"].isin(selected_types)]
    if selected_methods:
        result = result[result["payment_method"].isin(selected_methods)]
    if search:
        needle = upper(search)
        mask = result[["ra", "notes", "operator"]].fillna("").apply(
            lambda column: column.str.upper().str.contains(needle, regex=False)
        ).any(axis=1)
        result = result[mask]
    return result


def cash_dashboard(frame):
    st.header("Dashboard cassa")
    filtered = cash_filters(frame)
    if frame.empty:
        st.info("Nessun movimento di cassa presente.")
        return
    total_in = float(filtered.loc[filtered["effect"].gt(0), "effect"].sum()) if not filtered.empty else 0
    total_out = abs(float(filtered.loc[filtered["effect"].lt(0), "effect"].sum())) if not filtered.empty else 0
    current_balance = cash_balance(frame)
    open_deposits = frame[frame["ra"].fillna("").ne("")].groupby("ra", as_index=False)["effect"].sum()
    open_deposits = open_deposits[open_deposits["effect"].gt(0.009)]
    columns = st.columns(4)
    columns[0].metric("Saldo cassa attuale", f"€ {current_balance:,.2f}")
    columns[1].metric("Entrate filtrate", f"€ {total_in:,.2f}")
    columns[2].metric("Uscite filtrate", f"€ {total_out:,.2f}")
    columns[3].metric("RA con deposito aperto", len(open_deposits))
    if filtered.empty:
        st.info("Nessun movimento corrisponde ai filtri selezionati.")
        return
    daily = filtered.dropna(subset=["movement_date"]).copy()
    daily["giorno"] = daily["movement_date"].dt.date
    daily = daily.groupby("giorno", as_index=False).agg(entrate=("effect", lambda x: x[x > 0].sum()), uscite=("effect", lambda x: abs(x[x < 0].sum())))
    c1, c2 = st.columns(2)
    c1.plotly_chart(
        px.bar(daily, x="giorno", y=["entrate", "uscite"], barmode="group", title="Entrate e uscite giornaliere"),
        use_container_width=True,
    )
    by_type = filtered.groupby("movement_type", as_index=False)["amount"].sum()
    c2.plotly_chart(px.bar(by_type, x="movement_type", y="amount", title="Movimenti per tipologia"), use_container_width=True)
    if not open_deposits.empty:
        st.subheader("Depositi ancora presenti per RA")
        open_deposits = open_deposits.rename(columns={"ra": "RA", "effect": "Residuo"}).sort_values("Residuo", ascending=False)
        st.dataframe(open_deposits, use_container_width=True, hide_index=True, column_config={"Residuo": st.column_config.NumberColumn(format="€ %.2f")})


def cash_movement_form(frame):
    st.header("Inserimento e modifica movimento cassa")
    records = {row.id: row for row in frame.itertuples()} if not frame.empty else {}
    selected = st.selectbox(
        "Movimento da modificare (lascia vuoto per inserirne uno nuovo)",
        [""] + list(records),
        format_func=lambda item_id: "" if not item_id else (
            f"{records[item_id].movement_date.date()} | {records[item_id].movement_type} | "
            f"{records[item_id].ra or 'SENZA RA'} | € {records[item_id].amount:.2f}"
        ),
    )
    current = frame[frame["id"].eq(selected)].iloc[0].to_dict() if selected else {}
    prefix = current.get("id", "new_cash_movement")
    c1, c2, c3 = st.columns(3)
    movement_date = c1.date_input(
        "Data movimento *",
        value=current["movement_date"].date() if current and pd.notna(current.get("movement_date")) else date.today(),
        key=f"{prefix}_cash_date",
    )
    all_types = CASH_IN_TYPES + CASH_OUT_TYPES
    current_type = upper(current.get("movement_type", ""))
    movement_type = c2.selectbox(
        "Tipo movimento *", [""] + all_types,
        index=([""] + all_types).index(current_type) if current_type in all_types else 0,
        key=f"{prefix}_cash_type",
    )
    amount = c3.number_input(
        "Importo *", min_value=0.0, step=10.0,
        value=float(current.get("amount") or 0), key=f"{prefix}_cash_amount",
    )
    c4, c5, c6 = st.columns(3)
    ra = c4.text_input("RA", value=current.get("ra", ""), key=f"{prefix}_cash_ra")
    methods = ["CONTANTI", "ASSEGNO", "CARTA", "BONIFICO", "ALTRO"]
    current_method = upper(current.get("payment_method", "")) or "CONTANTI"
    method = c5.selectbox("Metodo *", methods, index=methods.index(current_method) if current_method in methods else 0, key=f"{prefix}_cash_method")
    linked_operator = current_operator() if not user_is_admin() else ""
    operator_options = [row["name"] for row in configured_operators()]
    if linked_operator and linked_operator not in operator_options:
        operator_options.append(linked_operator)
    current_operator_value = linked_operator or upper(current.get("operator", ""))
    if current_operator_value and current_operator_value not in operator_options:
        operator_options.append(current_operator_value)
    operator = c6.selectbox(
        "Operatore *", [""] + operator_options,
        index=([""] + operator_options).index(current_operator_value) if current_operator_value else 0,
        disabled=bool(linked_operator), key=f"{prefix}_cash_operator",
    )
    notes = st.text_area("Note", value=current.get("notes", ""), key=f"{prefix}_cash_notes")
    effect = cash_effect(movement_type, amount) if movement_type else 0
    st.info(f"Effetto sul saldo: € {effect:,.2f}")
    if st.button("Salva movimento", type="primary", use_container_width=True):
        if not movement_type or amount <= 0 or not operator:
            st.error("Compila tipo movimento, importo e operatore.")
        elif movement_type in {"DEPOSITO", "INCASSO", "RIMBORSO"} and not ra:
            st.error("Per depositi, incassi e rimborsi è obbligatorio indicare il RA.")
        else:
            save_cash_movement({
                "id": current.get("id"), "source_key": current.get("source_key"),
                "movement_date": movement_date.isoformat(), "ra": ra,
                "movement_type": movement_type, "amount": amount, "payment_method": method,
                "operator": operator, "notes": notes, "created_at": current.get("created_at"),
            })
            st.success("Movimento salvato.")
            st.rerun()
    if current:
        confirm = st.checkbox("Confermo l’eliminazione definitiva del movimento")
        if st.button("Elimina movimento", disabled=not confirm, use_container_width=True):
            with db() as conn:
                conn.execute("DELETE FROM cash_movements WHERE id = ?", (current["id"],))
            st.success("Movimento eliminato.")
            st.rerun()


def cash_archive(frame):
    st.header("Archivio movimenti cassa")
    filtered = cash_filters(frame)
    if filtered.empty:
        st.info("Nessun movimento disponibile.")
        return
    display = filtered[["movement_date", "ra", "movement_type", "amount", "effect", "payment_method", "operator", "notes"]].copy()
    display.columns = ["Data", "RA", "Tipo", "Importo", "Effetto saldo", "Metodo", "Operatore", "Note"]
    st.dataframe(
        display, use_container_width=True, hide_index=True,
        column_config={
            "Data": st.column_config.DateColumn(format="DD/MM/YYYY"),
            "Importo": st.column_config.NumberColumn(format="€ %.2f"),
            "Effetto saldo": st.column_config.NumberColumn(format="€ %.2f"),
        },
    )
    st.download_button(
        "Esporta archivio cassa in Excel", table_to_excel(display, "Movimenti cassa"),
        "archivio_cassa.xlsx", use_container_width=True,
    )


def cash_closing_page(frame):
    st.header("Conteggio e chiusura cassa")
    expected = cash_balance(frame)
    st.metric("Saldo contabile atteso", f"€ {expected:,.2f}")
    denominations = [200, 100, 50, 20, 10, 5, 2, 1, 0.50, 0.20, 0.10, 0.05]
    counts = {}
    st.subheader("Conteggio contanti")
    columns = st.columns(4)
    for index, denomination in enumerate(denominations):
        label = f"€ {denomination:g}"
        counts[str(denomination)] = columns[index % 4].number_input(label, min_value=0, step=1, key=f"cash_count_{denomination}")
    counted_cash = sum(float(value) * int(counts[str(value)]) for value in denominations)
    checks_total = st.number_input("Totale assegni presenti", min_value=0.0, step=10.0)
    physical_total = counted_cash + checks_total
    difference = physical_total - expected
    c1, c2, c3 = st.columns(3)
    c1.metric("Contanti contati", f"€ {counted_cash:,.2f}")
    c2.metric("Totale fisico", f"€ {physical_total:,.2f}")
    c3.metric("Differenza", f"€ {difference:,.2f}")
    operator = current_operator()
    if not operator:
        operator = st.selectbox("Operatore chiusura *", [""] + [row["name"] for row in configured_operators()])
    else:
        st.info(f"Operatore: {operator}")
    notes = st.text_area("Note chiusura")
    if st.button("Registra chiusura cassa", type="primary", use_container_width=True):
        if not operator:
            st.error("Seleziona l’operatore.")
        else:
            with db() as conn:
                conn.execute(
                    """
                    INSERT INTO cash_closings
                    (id, closing_date, expected_balance, counted_cash, checks_total, difference, denominations, operator, notes, created_at)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    ON CONFLICT(closing_date) DO UPDATE SET
                        expected_balance=excluded.expected_balance, counted_cash=excluded.counted_cash,
                        checks_total=excluded.checks_total, difference=excluded.difference,
                        denominations=excluded.denominations, operator=excluded.operator,
                        notes=excluded.notes, created_at=excluded.created_at
                    """,
                    (
                        str(uuid.uuid4()), date.today().isoformat(), expected, counted_cash, checks_total,
                        difference, json.dumps(counts), operator, clean(notes), datetime.now().isoformat(timespec="seconds"),
                    ),
                )
            st.success("Chiusura cassa registrata.")
            st.rerun()
    closings = load_cash_closings()
    if not closings.empty:
        st.subheader("Storico chiusure")
        display = closings[["closing_date", "expected_balance", "counted_cash", "checks_total", "difference", "operator", "notes"]].copy()
        display.columns = ["Data", "Saldo atteso", "Contanti", "Assegni", "Differenza", "Operatore", "Note"]
        st.dataframe(display, use_container_width=True, hide_index=True)


def import_cash_workbook(file_bytes):
    book = pd.ExcelFile(io.BytesIO(file_bytes))
    added = 0
    skipped = 0
    opening_added = False
    for sheet in book.sheet_names:
        raw = pd.read_excel(io.BytesIO(file_bytes), sheet_name=sheet, header=None, dtype=object)
        header_row = None
        for index in range(min(12, len(raw))):
            values = [upper(value) for value in raw.iloc[index, :6].tolist()]
            if "RA" in values and any("DATA APERTURA" in value for value in values):
                header_row = index
                break
        if header_row is None:
            continue
        if not opening_added and header_row + 1 < len(raw):
            opening_label = upper(raw.iloc[header_row + 1, 0])
            opening_value = pd.to_numeric(raw.iloc[header_row + 1, 5], errors="coerce")
            if "RIMANENZA" in opening_label and pd.notna(opening_value) and float(opening_value) != 0:
                source_dates = pd.concat([
                    pd.to_datetime(raw.iloc[header_row + 1:, 1], errors="coerce"),
                    pd.to_datetime(raw.iloc[header_row + 1:, 3], errors="coerce"),
                ]).dropna()
                opening_date = source_dates.min().date() if not source_dates.empty else date.today()
                save_cash_movement({
                    "source_key": f"cash-opening|{sheet}|{header_row + 2}",
                    "movement_date": opening_date.isoformat(), "movement_type": "RETTIFICA POSITIVA",
                    "amount": abs(float(opening_value)), "payment_method": "CONTANTI",
                    "operator": "IMPORTAZIONE", "notes": clean(raw.iloc[header_row + 1, 0]),
                })
                opening_added = True
                added += 1
        for row_index in range(header_row + 1, len(raw)):
            row = raw.iloc[row_index]
            ra_raw = clean(row.iloc[0] if len(row) > 0 else "")
            if not ra_raw or "RIMANENZA" in upper(ra_raw) or "TOTALE CASSA" in upper(ra_raw):
                continue
            deposit = pd.to_numeric(row.iloc[2] if len(row) > 2 else None, errors="coerce")
            refund = pd.to_numeric(row.iloc[4] if len(row) > 4 else None, errors="coerce")
            open_date = pd.to_datetime(row.iloc[1] if len(row) > 1 else None, errors="coerce")
            close_date = pd.to_datetime(row.iloc[3] if len(row) > 3 else None, errors="coerce")
            is_remittance = "RIMESSA" in upper(ra_raw)
            imported_row = False
            if pd.notna(deposit) and float(deposit) != 0 and pd.notna(open_date):
                save_cash_movement({
                    "source_key": f"cash|{sheet}|{row_index + 1}|in",
                    "movement_date": open_date.date().isoformat(), "ra": "" if is_remittance else ra_raw,
                    "movement_type": "DEPOSITO", "amount": abs(float(deposit)),
                    "payment_method": "CONTANTI", "operator": "IMPORTAZIONE",
                    "notes": f"Storico Excel — {sheet}",
                })
                added += 1
                imported_row = True
            if pd.notna(refund) and float(refund) != 0 and pd.notna(close_date):
                save_cash_movement({
                    "source_key": f"cash|{sheet}|{row_index + 1}|out",
                    "movement_date": close_date.date().isoformat(), "ra": "" if is_remittance else ra_raw,
                    "movement_type": "RIMESSA" if is_remittance else "RIMBORSO",
                    "amount": abs(float(refund)), "payment_method": "CONTANTI",
                    "operator": "IMPORTAZIONE", "notes": f"Storico Excel — {sheet}",
                })
                added += 1
                imported_row = True
            if not imported_row and (pd.notna(deposit) or pd.notna(refund)):
                skipped += 1
    return added, skipped, len(book.sheet_names)


def cash_import_page():
    st.header("Importazione storico cassa")
    st.caption("Importa i fogli mensili con colonne RA, data apertura, deposito/incasso, data chiusura e rimborso.")
    upload = st.file_uploader("File Excel cassa", type=["xlsx"], key="cash_history_upload")
    if upload and st.button("Importa storico cassa", type="primary", use_container_width=True):
        added, skipped, sheets = import_cash_workbook(upload.getvalue())
        st.success(f"Elaborati {sheets} fogli e importati o aggiornati {added} movimenti.")
        if skipped:
            st.info(f"Righe non importabili perché prive di data: {skipped}.")
        st.rerun()


def special_events_dashboard(frame):
    st.header("Dashboard eventi speciali")
    if frame.empty:
        st.info("Nessun evento presente. Usa Crea evento per inserire il primo.")
        return
    today = pd.Timestamp(date.today())
    upcoming = frame[frame["start_date"].ge(today) & ~frame["status"].isin(["ANNULLATO", "CONCLUSO"])]
    columns = st.columns(6)
    columns[0].metric("Eventi totali", len(frame))
    columns[1].metric("In programma", len(upcoming))
    columns[2].metric("Partecipanti", int(frame["participants"].sum()))
    columns[3].metric("Ricavi", f"€ {frame['revenue'].sum():,.2f}")
    columns[4].metric("Margine", f"€ {frame['margin'].sum():,.2f}")
    columns[5].metric("Veicoli assegnati", int(frame["assigned_vehicles"].sum()))
    by_status = frame.groupby("status", as_index=False).size().rename(columns={"size": "eventi"})
    by_type = frame.groupby("event_type", as_index=False).agg(eventi=("id", "count"), ricavi=("revenue", "sum"))
    c1, c2 = st.columns(2)
    c1.plotly_chart(px.bar(by_status, x="status", y="eventi", title="Eventi per stato"), use_container_width=True)
    c2.plotly_chart(px.bar(by_type, x="event_type", y="eventi", color="ricavi", title="Eventi per tipologia"), use_container_width=True)
    st.subheader("Prossimi eventi")
    if upcoming.empty:
        st.info("Nessun evento futuro programmato.")
    else:
        display = upcoming[["start_date", "title", "event_type", "location", "responsible", "status", "participants"]].copy()
        display.columns = ["Data", "Evento", "Tipologia", "Luogo", "Responsabile", "Stato", "Partecipanti"]
        st.dataframe(display.sort_values("Data"), use_container_width=True, hide_index=True)


def special_event_form(frame):
    st.header("Crea o modifica evento")
    records, record_ids = event_options(frame)
    selected = st.selectbox(
        "Evento da modificare (lascia vuoto per crearne uno nuovo)", [""] + record_ids,
        format_func=lambda item_id: "" if not item_id else f"{records[item_id].start_date.date()} | {records[item_id].title}",
    )
    current = frame[frame["id"].eq(selected)].iloc[0].to_dict() if selected else {}
    prefix = current.get("id", "new_special_event")
    c1, c2, c3 = st.columns(3)
    title = c1.text_input("Nome evento *", value=current.get("title", ""), key=f"{prefix}_event_title")
    event_types = ["PROMOZIONALE", "AZIENDALE", "FIERA", "CONSEGNA SPECIALE", "ESPOSIZIONE", "ALTRO"]
    current_type = upper(current.get("event_type", ""))
    if current_type and current_type not in event_types:
        event_types.append(current_type)
    event_type = c2.selectbox("Tipologia *", [""] + event_types, index=([""] + event_types).index(current_type) if current_type else 0, key=f"{prefix}_event_type")
    statuses = ["PROGRAMMATO", "CONFERMATO", "IN CORSO", "CONCLUSO", "ANNULLATO"]
    current_status = upper(current.get("status", "")) or "PROGRAMMATO"
    status = c3.selectbox("Stato", statuses, index=statuses.index(current_status), key=f"{prefix}_event_status")
    c4, c5, c6, c7 = st.columns(4)
    start_date_value = current.get("start_date")
    end_date_value = current.get("end_date")
    start_date = c4.date_input("Data inizio *", value=start_date_value.date() if pd.notna(start_date_value) else date.today(), key=f"{prefix}_event_start_date")
    end_date = c5.date_input("Data fine", value=end_date_value.date() if pd.notna(end_date_value) else start_date, key=f"{prefix}_event_end_date")
    start_time = c6.text_input("Ora inizio", value=current.get("start_time", ""), placeholder="09:00", key=f"{prefix}_event_start_time")
    end_time = c7.text_input("Ora fine", value=current.get("end_time", ""), placeholder="18:00", key=f"{prefix}_event_end_time")
    c8, c9, c10 = st.columns(3)
    location = c8.text_input("Luogo *", value=current.get("location", ""), key=f"{prefix}_event_location")
    responsible_options = [row["name"] for row in configured_operators()]
    current_responsible = upper(current.get("responsible", ""))
    if current_responsible and current_responsible not in responsible_options:
        responsible_options.append(current_responsible)
    responsible = c9.selectbox("Responsabile", [""] + responsible_options, index=([""] + responsible_options).index(current_responsible) if current_responsible else 0, key=f"{prefix}_event_responsible")
    capacity = c10.number_input("Capienza prevista", min_value=0, step=1, value=int(current.get("capacity") or 0), key=f"{prefix}_event_capacity")
    vehicles = st.text_input("Note generali veicoli", value=current.get("vehicles", ""), placeholder="L’elenco dettagliato si gestisce nella pagina Veicoli evento", key=f"{prefix}_event_vehicles")
    c11, c12 = st.columns(2)
    cost = c11.number_input("Costi previsti", min_value=0.0, step=50.0, value=float(current.get("cost") or 0), key=f"{prefix}_event_cost")
    revenue = c12.number_input("Incassi previsti", min_value=0.0, step=50.0, value=float(current.get("revenue") or 0), key=f"{prefix}_event_revenue")
    description = st.text_area("Descrizione", value=current.get("description", ""), key=f"{prefix}_event_description")
    notes = st.text_area("Note", value=current.get("notes", ""), key=f"{prefix}_event_notes")
    st.info(f"Margine previsto: € {revenue - cost:,.2f}")
    if st.button("Salva evento", type="primary", use_container_width=True):
        if not all([title, event_type, start_date, location]):
            st.error("Compila nome, tipologia, data iniziale e luogo.")
        elif end_date < start_date:
            st.error("La data finale non può precedere quella iniziale.")
        else:
            save_special_event({
                "id": current.get("id"), "title": title, "event_type": event_type,
                "start_date": start_date.isoformat(), "end_date": end_date.isoformat(),
                "start_time": start_time, "end_time": end_time, "location": location,
                "responsible": responsible, "status": status, "capacity": capacity,
                "vehicles": vehicles, "cost": cost, "revenue": revenue,
                "description": description, "notes": notes,
                "created_by": current.get("created_by"), "created_at": current.get("created_at"),
            })
            st.success("Evento salvato.")
            st.rerun()
    if current:
        confirm = st.checkbox("Confermo l’eliminazione definitiva dell’evento e dei relativi allegati")
        if st.button("Elimina evento", disabled=not confirm, use_container_width=True):
            with db() as conn:
                conn.execute("DELETE FROM event_participants WHERE event_id = ?", (current["id"],))
                conn.execute("DELETE FROM event_files WHERE event_id = ?", (current["id"],))
                conn.execute("DELETE FROM event_vehicles WHERE event_id = ?", (current["id"],))
                conn.execute("DELETE FROM special_events WHERE id = ?", (current["id"],))
            st.success("Evento eliminato.")
            st.rerun()


def events_calendar(frame):
    st.header("Calendario eventi")
    if frame.empty:
        st.info("Nessun evento presente.")
        return
    years = sorted(frame["start_date"].dropna().dt.year.unique().astype(int).tolist())
    selected_year = st.selectbox("Anno", years, index=len(years) - 1)
    result = frame[frame["start_date"].dt.year.eq(selected_year)].copy()
    result["Mese"] = result["start_date"].dt.strftime("%Y-%m")
    display = result[["start_date", "end_date", "title", "event_type", "location", "start_time", "end_time", "responsible", "status"]].copy()
    display.columns = ["Data inizio", "Data fine", "Evento", "Tipologia", "Luogo", "Ora inizio", "Ora fine", "Responsabile", "Stato"]
    st.dataframe(display.sort_values("Data inizio"), use_container_width=True, hide_index=True)


def events_archive(frame):
    st.header("Archivio eventi speciali")
    if frame.empty:
        st.info("Nessun evento presente.")
        return
    c1, c2, c3 = st.columns(3)
    selected_status = c1.multiselect("Stato", sorted(frame["status"].dropna().unique()))
    selected_types = c2.multiselect("Tipologia", sorted(frame["event_type"].dropna().unique()))
    search = c3.text_input("Cerca evento o luogo")
    result = frame.copy()
    if selected_status:
        result = result[result["status"].isin(selected_status)]
    if selected_types:
        result = result[result["event_type"].isin(selected_types)]
    if search:
        needle = upper(search)
        result = result[result[["title", "location", "description"]].fillna("").apply(lambda col: col.str.upper().str.contains(needle, regex=False)).any(axis=1)]
    display = result[["start_date", "title", "event_type", "location", "responsible", "status", "participants", "vehicles", "cost", "revenue", "margin", "files"]].copy()
    display.columns = ["Data", "Evento", "Tipologia", "Luogo", "Responsabile", "Stato", "Partecipanti", "Veicoli", "Costi", "Incassi", "Margine", "Allegati"]
    st.dataframe(display, use_container_width=True, hide_index=True)
    st.download_button("Esporta archivio eventi", table_to_excel(display, "Eventi speciali"), "eventi_speciali.xlsx", use_container_width=True)


def event_participants_page(frame):
    st.header("Partecipanti e personale")
    records, record_ids = event_options(frame)
    if not record_ids:
        st.info("Crea prima un evento.")
        return
    event_id = st.selectbox("Evento", record_ids, format_func=lambda item_id: f"{records[item_id].start_date.date()} | {records[item_id].title}")
    with st.form("add_event_participant", clear_on_submit=True):
        c1, c2, c3 = st.columns(3)
        name = c1.text_input("Nome *")
        role = c2.selectbox("Ruolo", ["PARTECIPANTE", "OPERATORE", "RESPONSABILE", "FORNITORE", "OSPITE", "ALTRO"])
        contact = c3.text_input("Contatto")
        notes = st.text_input("Note")
        add = st.form_submit_button("Aggiungi", type="primary", use_container_width=True)
    if add:
        if not name:
            st.error("Inserisci il nome.")
        else:
            with db() as conn:
                conn.execute("INSERT INTO event_participants VALUES (?, ?, ?, ?, ?, ?, ?)", (str(uuid.uuid4()), event_id, clean(name), upper(role), clean(contact), clean(notes), datetime.now().isoformat(timespec="seconds")))
            st.rerun()
    with db() as conn:
        participants = [dict(row) for row in conn.execute("SELECT * FROM event_participants WHERE event_id = ? ORDER BY role, name", (event_id,)).fetchall()]
    if participants:
        st.dataframe(pd.DataFrame(participants)[["name", "role", "contact", "notes"]], use_container_width=True, hide_index=True)
        remove_id = st.selectbox("Persona da rimuovere", [row["id"] for row in participants], format_func=lambda item_id: next(row["name"] for row in participants if row["id"] == item_id))
        if st.button("Rimuovi persona", use_container_width=True):
            with db() as conn:
                conn.execute("DELETE FROM event_participants WHERE id = ?", (remove_id,))
            st.rerun()
    else:
        st.info("Nessuna persona associata.")


def event_files_page(frame):
    st.header("Documenti e fotografie evento")
    records, record_ids = event_options(frame)
    if not record_ids:
        st.info("Crea prima un evento.")
        return
    event_id = st.selectbox("Evento", record_ids, format_func=lambda item_id: f"{records[item_id].start_date.date()} | {records[item_id].title}")
    uploads = st.file_uploader("Carica documenti o fotografie", accept_multiple_files=True, key=f"event_upload_{event_id}")
    if st.button("Archivia allegati", type="primary", disabled=not uploads, use_container_width=True):
        rows = [(str(uuid.uuid4()), event_id, clean(file.name), clean(file.type), sqlite3.Binary(file.getvalue()), datetime.now().isoformat(timespec="seconds")) for file in uploads if file.getvalue()]
        with db() as conn:
            conn.executemany("INSERT INTO event_files VALUES (?, ?, ?, ?, ?, ?)", rows)
        st.success(f"Allegati archiviati: {len(rows)}.")
        st.rerun()
    with db() as conn:
        files = [dict(row) for row in conn.execute("SELECT * FROM event_files WHERE event_id = ? ORDER BY created_at DESC", (event_id,)).fetchall()]
    if not files:
        st.info("Nessun allegato presente.")
        return
    for file in files:
        c1, c2, c3 = st.columns([4, 1, 1])
        c1.write(file["file_name"])
        c2.download_button("Scarica", file["file_data"], file["file_name"], mime=file["mime_type"] or "application/octet-stream", key=f"event_download_{file['id']}")
        if c3.button("Elimina", key=f"event_delete_{file['id']}"):
            with db() as conn:
                conn.execute("DELETE FROM event_files WHERE id = ?", (file["id"],))
            st.rerun()


def event_vehicles_page(frame):
    st.header("Veicoli assegnati all’evento")
    st.caption(
        "Inserisci o incolla l’elenco direttamente nella tabella. Le righe vengono ordinate per gruppo, "
        "targa, marca e modello. Il pallino diventa verde quando è indicata la data di ritiro."
    )
    records, record_ids = event_options(frame)
    if not record_ids:
        st.info("Crea prima un evento.")
        return
    event_id = st.selectbox(
        "Evento",
        record_ids,
        format_func=lambda item_id: f"{records[item_id].start_date.date()} | {records[item_id].title}",
        key="event_vehicle_event",
    )
    with db() as conn:
        stored = [
            dict(row) for row in conn.execute(
                """
                SELECT * FROM event_vehicles
                WHERE event_id = ?
                ORDER BY UPPER(COALESCE(vehicle_group, '')), UPPER(COALESCE(plate, '')),
                         UPPER(COALESCE(brand, '')), UPPER(COALESCE(model, ''))
                """,
                (event_id,),
            ).fetchall()
        ]
    st.subheader("Inserimento o modifica veicolo")
    vehicle_by_id = {item["id"]: item for item in stored}
    selected_vehicle_id = st.selectbox(
        "Veicolo da modificare",
        [""] + list(vehicle_by_id),
        format_func=lambda item_id: "➕ NUOVO VEICOLO" if not item_id else " | ".join(
            value for value in [
                upper(vehicle_by_id[item_id].get("plate")),
                clean(vehicle_by_id[item_id].get("brand")),
                clean(vehicle_by_id[item_id].get("model")),
                clean(vehicle_by_id[item_id].get("assigned_to")),
            ] if value
        ),
        key=f"event_vehicle_selector_{event_id}",
    )
    selected_vehicle = vehicle_by_id.get(selected_vehicle_id, {})
    form_key = selected_vehicle_id or "new"
    existing_pickup = clean(selected_vehicle.get("pickup_date"))
    pickup_default = date.today()
    if existing_pickup:
        parsed_pickup = pd.to_datetime(existing_pickup, errors="coerce", dayfirst=True)
        if pd.notna(parsed_pickup):
            pickup_default = parsed_pickup.date()
    with st.form(f"event_vehicle_form_{event_id}_{form_key}", clear_on_submit=not bool(selected_vehicle_id)):
        c1, c2, c3, c4 = st.columns(4)
        vehicle_group = c1.text_input("Gruppo", value=clean(selected_vehicle.get("vehicle_group")))
        plate = c2.text_input("Targa *", value=upper(selected_vehicle.get("plate")))
        brand = c3.text_input("Marca", value=clean(selected_vehicle.get("brand")))
        model = c4.text_input("Modello", value=clean(selected_vehicle.get("model")))
        c5, c6, c7 = st.columns(3)
        assigned_to = c5.text_input("Assegnata a", value=clean(selected_vehicle.get("assigned_to")))
        ra = c6.text_input("RA", value=upper(selected_vehicle.get("ra")))
        picked_up = c7.checkbox("Veicolo ritirato", value=bool(existing_pickup))
        c8, c9 = st.columns([1, 2])
        pickup_date = c8.date_input("Data ritiro", value=pickup_default, disabled=not picked_up)
        vehicle_notes = c9.text_input("Note", value=clean(selected_vehicle.get("notes")))
        save_vehicle = st.form_submit_button(
            "Aggiorna veicolo" if selected_vehicle_id else "Aggiungi veicolo alla tabella",
            type="primary", use_container_width=True,
        )
    if save_vehicle:
        if not clean(plate):
            st.error("Inserisci almeno la targa del veicolo.")
        else:
            vehicle_id = selected_vehicle_id or str(uuid.uuid4())
            with db() as conn:
                conn.execute(
                    """
                    INSERT INTO event_vehicles
                    (id,event_id,vehicle_group,plate,brand,model,assigned_to,ra,pickup_date,notes,updated_at)
                    VALUES (?,?,?,?,?,?,?,?,?,?,?)
                    ON CONFLICT(id) DO UPDATE SET
                        vehicle_group=excluded.vehicle_group,plate=excluded.plate,brand=excluded.brand,
                        model=excluded.model,assigned_to=excluded.assigned_to,ra=excluded.ra,
                        pickup_date=excluded.pickup_date,notes=excluded.notes,updated_at=excluded.updated_at
                    """,
                    (
                        vehicle_id, event_id, clean(vehicle_group), upper(plate), clean(brand), clean(model),
                        clean(assigned_to), upper(ra), pickup_date.isoformat() if picked_up else "",
                        clean(vehicle_notes), datetime.now().isoformat(timespec="seconds"),
                    ),
                )
            st.success("Veicolo aggiornato nella tabella." if selected_vehicle_id else "Veicolo aggiunto automaticamente alla tabella.")
            st.rerun()
    if selected_vehicle_id:
        confirm_vehicle_delete = st.checkbox(
            "Confermo l’eliminazione del veicolo selezionato",
            key=f"confirm_vehicle_delete_{event_id}_{selected_vehicle_id}",
        )
        if st.button(
            "Elimina veicolo selezionato", disabled=not confirm_vehicle_delete,
            key=f"delete_vehicle_{event_id}_{selected_vehicle_id}", use_container_width=True,
        ):
            with db() as conn:
                conn.execute("DELETE FROM event_vehicles WHERE id = ? AND event_id = ?", (selected_vehicle_id, event_id))
            st.success("Veicolo eliminato.")
            st.rerun()
    st.divider()
    st.subheader("Elenco veicoli dell’evento")
    columns = ["_id", "Stato", "Gruppo", "Targa", "Marca", "Modello", "Assegnata a", "RA", "Data ritiro", "Note"]
    rows = []
    for item in stored:
        pickup = clean(item.get("pickup_date"))
        rows.append({
            "_id": item["id"], "Stato": "🟢" if pickup else "🔴",
            "Gruppo": item.get("vehicle_group", ""), "Targa": item.get("plate", ""),
            "Marca": item.get("brand", ""), "Modello": item.get("model", ""),
            "Assegnata a": item.get("assigned_to", ""), "RA": item.get("ra", ""),
            "Data ritiro": pickup, "Note": item.get("notes", ""),
        })
    source = pd.DataFrame(rows, columns=columns)
    if source.empty:
        source = pd.DataFrame(columns=columns)
    total = len(stored)
    collected = sum(bool(clean(item.get("pickup_date"))) for item in stored)
    m1, m2, m3 = st.columns(3)
    m1.metric("Veicoli assegnati", total, delta=f"su 40 previsti" if total <= 40 else "oltre 40 previsti")
    m2.metric("Ritirati", collected)
    m3.metric("Da ritirare", max(total - collected, 0))
    edited = st.data_editor(
        source,
        num_rows="dynamic",
        use_container_width=True,
        hide_index=True,
        key=f"event_vehicle_editor_{event_id}",
        column_config={
            "_id": None,
            "Stato": st.column_config.TextColumn("Ritiro", disabled=True, width="small", help="🟢 ritirata; 🔴 non ancora ritirata"),
            "Gruppo": st.column_config.TextColumn("Gruppo", width="medium"),
            "Targa": st.column_config.TextColumn("Targa", width="small"),
            "Marca": st.column_config.TextColumn("Marca", width="medium"),
            "Modello": st.column_config.TextColumn("Modello", width="medium"),
            "Assegnata a": st.column_config.TextColumn("Assegnata a", width="medium"),
            "RA": st.column_config.TextColumn("RA", width="small"),
            "Data ritiro": st.column_config.TextColumn("Data ritiro", help="Esempio: 25/09/2026", width="small"),
            "Note": st.column_config.TextColumn("Note", width="large"),
        },
        disabled=["Stato"],
    )
    st.caption("Puoi copiare più righe da Excel e incollarle nella prima cella. Usa il + in fondo per aggiungere una riga.")
    export = edited.copy()
    if "_id" in export.columns:
        export = export.drop(columns=["_id"])
    export["Stato"] = export.get("Data ritiro", pd.Series(dtype=str)).fillna("").astype(str).str.strip().map(
        lambda value: "RITIRATA" if value else "DA RITIRARE"
    )
    export = export.sort_values(
        ["Gruppo", "Targa", "Marca", "Modello"],
        key=lambda column: column.fillna("").astype(str).str.upper(),
    )
    safe_event_name = re.sub(r"[^A-Za-z0-9_-]+", "_", str(records[event_id].title)).strip("_") or "evento"
    d1, d2 = st.columns(2)
    d1.download_button(
        "Scarica elenco in Excel", table_to_excel(export, "Veicoli evento"),
        f"veicoli_{safe_event_name}.xlsx", use_container_width=True,
    )
    d2.download_button(
        "Scarica elenco in PDF", event_vehicles_to_pdf(export, str(records[event_id].title)),
        f"veicoli_{safe_event_name}.pdf", mime="application/pdf", use_container_width=True,
    )
    if st.button("Salva elenco veicoli", type="primary", use_container_width=True):
        saved_ids = []
        now = datetime.now().isoformat(timespec="seconds")
        with db() as conn:
            for _, row in edited.fillna("").iterrows():
                values = {name: clean(row.get(name, "")) for name in ["Gruppo", "Targa", "Marca", "Modello", "Assegnata a", "RA", "Data ritiro", "Note"]}
                if not any(values.values()):
                    continue
                vehicle_id = clean(row.get("_id")) or str(uuid.uuid4())
                saved_ids.append(vehicle_id)
                conn.execute(
                    """
                    INSERT INTO event_vehicles
                    (id,event_id,vehicle_group,plate,brand,model,assigned_to,ra,pickup_date,notes,updated_at)
                    VALUES (?,?,?,?,?,?,?,?,?,?,?)
                    ON CONFLICT(id) DO UPDATE SET
                        vehicle_group=excluded.vehicle_group,plate=excluded.plate,brand=excluded.brand,
                        model=excluded.model,assigned_to=excluded.assigned_to,ra=excluded.ra,
                        pickup_date=excluded.pickup_date,notes=excluded.notes,updated_at=excluded.updated_at
                    """,
                    (vehicle_id, event_id, values["Gruppo"], upper(values["Targa"]), values["Marca"],
                     values["Modello"], values["Assegnata a"], upper(values["RA"]),
                     values["Data ritiro"], values["Note"], now),
                )
            if saved_ids:
                placeholders = ",".join("?" for _ in saved_ids)
                conn.execute(
                    f"DELETE FROM event_vehicles WHERE event_id = ? AND id NOT IN ({placeholders})",
                    [event_id] + saved_ids,
                )
            else:
                conn.execute("DELETE FROM event_vehicles WHERE event_id = ?", (event_id,))
        st.success(f"Elenco salvato: {len(saved_ids)} veicoli.")
        st.rerun()


def import_backup():
    st.header("Importazione e backup")
    st.subheader("Backup completo")
    if DB_PATH.exists():
        st.download_button("Scarica database", DB_PATH.read_bytes(), "ancillary_backup.db", mime="application/octet-stream", use_container_width=True)
    st.caption("Conserva periodicamente il file di backup in una posizione sicura.")
    st.subheader("Importazione intelligente da Excel")
    st.caption("Il gestionale riconosce automaticamente riepiloghi ancillary, report Analisi Contratti RA e file Addebito Danni.")
    upload = st.file_uploader("Carica un file Excel", type=["xlsx"])
    if upload:
        file_bytes = upload.getvalue()
        detected = detect_excel_type(file_bytes)
        labels = {
            "ANCILLARY": "Riepilogo ancillary", "CONTRATTI_RA": "Analisi Contratti RA",
            "ADDEBITO_DANNI": "Riepilogo addebito danni", "SCONOSCIUTO": "Formato non riconosciuto",
        }
        st.info(f"Formato rilevato: **{labels[detected]}**")
        if st.button("Importa, scorpora e analizza", type="primary", disabled=detected == "SCONOSCIUTO", use_container_width=True):
            if detected == "CONTRATTI_RA":
                added, main_sheet, period = import_contract_workbook(file_bytes, upload.name)
                st.success(f"Importati o aggiornati {added} contratti del periodo {period}. Foglio principale: {main_sheet}.")
                st.rerun()
            if detected == "ADDEBITO_DANNI":
                added, skipped = import_damage_workbook(file_bytes, upload.name)
                st.success(f"Importate o aggiornate {added} segnalazioni danni.")
                if skipped:
                    st.info(f"Escluse {skipped} righe senza numero RA.")
                st.rerun()
            imported = pd.read_excel(io.BytesIO(file_bytes), sheet_name=0, dtype=object)
            added = 0
            skipped_before_start = 0
            for _, row in imported.iterrows():
                days = pd.to_numeric(row.get("GIORNI NOLEGGIO"), errors="coerce")
                cost = pd.to_numeric(row.get("COSTO TOTALE ANCILLARY (iva esclusa)"), errors="coerce")
                start = pd.to_datetime(row.get("DATA INIZIO NOLEGGIO"), errors="coerce")
                if pd.notna(start) and start.date() < ANCILLARY_START_DATE:
                    skipped_before_start += 1
                    continue
                vehicle = upper(row.get("TIPO DI VEICOLO"))
                ancillary = upper(row.get("SCELTA ANCILLARY (CAR)")) or upper(row.get("SCELTA ANCILLARY (VAN)"))
                save_record({
                    "ra": row.get("RA (Rental Agreement)"),
                    "start_date": "" if pd.isna(start) else start.date().isoformat(),
                    "rental_days": 0 if pd.isna(days) else int(days),
                    "source_raw": row.get("FONTE"), "operator": row.get("OPERATORE"),
                    "vehicle_type": vehicle, "ancillary": ancillary,
                    "ancillary_cost": None if pd.isna(cost) else float(cost),
                    "rental_type": row.get("TIPO NOLEGGIO"), "notes": "Importato da Excel",
                })
                added += 1
            st.success(f"Importati {added} record ancillary.")
            if skipped_before_start:
                st.info(f"Esclusi {skipped_before_start} record anteriori al 01/10/2026.")
            st.rerun()


def operator_settings():
    st.header("Configurazione operatori")
    st.caption("Gli operatori attivi compaiono nella tendina del modulo di inserimento. Lo storico non viene modificato.")

    with st.form("add_operator", clear_on_submit=True):
        new_operator = st.text_input("Nuovo operatore")
        add_operator = st.form_submit_button("Aggiungi operatore", type="primary", use_container_width=True)
    if add_operator:
        name = upper(new_operator)
        if not name:
            st.error("Inserisci il nome dell’operatore.")
        else:
            with db() as conn:
                conn.execute(
                    "INSERT INTO operator_config (name, active) VALUES (?, 1) "
                    "ON CONFLICT(name) DO UPDATE SET active = 1",
                    (name,),
                )
            st.success(f"Operatore {name} disponibile nel modulo.")
            st.rerun()

    rows = configured_operators(active_only=False)
    if not rows:
        st.info("Nessun operatore configurato.")
        return

    st.subheader("Operatori configurati")
    st.dataframe(
        pd.DataFrame(rows).rename(columns={"name": "Operatore", "active": "Attivo"}),
        use_container_width=True,
        hide_index=True,
        column_config={"Attivo": st.column_config.CheckboxColumn()},
    )
    selected_name = st.selectbox("Operatore da gestire", [row["name"] for row in rows])
    selected_row = next(row for row in rows if row["name"] == selected_name)
    c1, c2 = st.columns(2)
    renamed = c1.text_input("Nuovo nome", value=selected_name, key=f"rename_{selected_name}")
    active = c2.checkbox("Attivo nella tendina", value=bool(selected_row["active"]), key=f"active_{selected_name}")
    if st.button("Salva modifiche operatore", type="primary", use_container_width=True):
        new_name = upper(renamed)
        if not new_name:
            st.error("Il nome non può essere vuoto.")
        else:
            try:
                with db() as conn:
                    conn.execute(
                        "UPDATE operator_config SET name = ?, active = ? WHERE name = ?",
                        (new_name, int(active), selected_name),
                    )
                    if new_name != selected_name:
                        conn.execute("UPDATE rentals SET operator = ? WHERE UPPER(TRIM(operator)) = ?", (new_name, selected_name))
                        conn.execute("UPDATE contracts SET operator = ? WHERE UPPER(TRIM(operator)) = ?", (new_name, selected_name))
                        conn.execute("UPDATE damage_charges SET operator = ? WHERE UPPER(TRIM(operator)) = ?", (new_name, selected_name))
                st.success("Operatore aggiornato.")
                st.rerun()
            except sqlite3.IntegrityError:
                st.error("Esiste già un operatore con questo nome.")

    confirm_delete = st.checkbox("Confermo la rimozione dalla configurazione")
    if st.button("Elimina operatore", disabled=not confirm_delete, use_container_width=True):
        with db() as conn:
            conn.execute("DELETE FROM operator_config WHERE name = ?", (selected_name,))
        st.success("Operatore rimosso dalla tendina. I dati storici restano invariati.")
        st.rerun()


def user_settings():
    st.header("Utenti e autorizzazioni")
    if not user_is_admin():
        st.error("Questa funzione è riservata agli amministratori.")
        return

    st.caption("Crea gli accessi personali, collega l’operatore e scegli le aree visibili nel menu.")
    operator_names = [row["name"] for row in configured_operators()]

    with st.expander("Crea un nuovo utente", expanded=True):
        with st.form("create_app_user", clear_on_submit=True):
            c1, c2 = st.columns(2)
            username = c1.text_input("Username *")
            display_name = c2.text_input("Nome visualizzato *")
            c3, c4 = st.columns(2)
            password = c3.text_input("Password iniziale *", type="password")
            operator_name = c4.selectbox("Operatore collegato", [""] + operator_names)
            permissions = st.multiselect("Aree autorizzate *", PERMISSION_AREAS)
            is_admin = st.checkbox("Amministratore: accesso completo e gestione utenti")
            submitted = st.form_submit_button("Crea utente", type="primary", use_container_width=True)
        if submitted:
            normalized_username = clean(username).lower()
            if not normalized_username or not display_name:
                st.error("Inserisci username e nome visualizzato.")
            elif len(password) < 8:
                st.error("La password deve contenere almeno 8 caratteri.")
            elif not is_admin and not permissions:
                st.error("Seleziona almeno un’area autorizzata.")
            else:
                try:
                    with db() as conn:
                        conn.execute(
                            """
                            INSERT INTO app_users
                            (username, password_hash, display_name, operator_name, permissions, active, is_admin, created_at)
                            VALUES (?, ?, ?, ?, ?, 1, ?, ?)
                            """,
                            (
                                normalized_username, password_hash(password), clean(display_name),
                                upper(operator_name), json.dumps(permissions), int(is_admin),
                                datetime.now().isoformat(timespec="seconds"),
                            ),
                        )
                    st.success(f"Utente {normalized_username} creato.")
                    st.rerun()
                except sqlite3.IntegrityError:
                    st.error("Questo username esiste già.")

    with db() as conn:
        users = [dict(row) for row in conn.execute("SELECT * FROM app_users ORDER BY username").fetchall()]
    if not users:
        st.info("Non sono ancora presenti utenti personali. L’accesso admin configurato nei Secrets resta attivo.")
        return

    summary = pd.DataFrame([
        {
            "Username": row["username"], "Nome": row["display_name"],
            "Operatore": row["operator_name"], "Aree": " · ".join(json.loads(row["permissions"] or "[]")),
            "Amministratore": bool(row["is_admin"]), "Attivo": bool(row["active"]),
        }
        for row in users
    ])
    st.subheader("Utenti configurati")
    st.dataframe(summary, use_container_width=True, hide_index=True)

    selected_username = st.selectbox("Utente da modificare", [row["username"] for row in users])
    selected = next(row for row in users if row["username"] == selected_username)
    selected_permissions = json.loads(selected["permissions"] or "[]")
    with st.form(f"edit_user_{selected_username}"):
        c1, c2 = st.columns(2)
        edited_name = c1.text_input("Nome visualizzato", value=selected["display_name"] or "")
        existing_operator = upper(selected["operator_name"])
        edit_operator_options = [""] + operator_names
        if existing_operator and existing_operator not in edit_operator_options:
            edit_operator_options.append(existing_operator)
        edited_operator = c2.selectbox(
            "Operatore collegato", edit_operator_options,
            index=edit_operator_options.index(existing_operator) if existing_operator in edit_operator_options else 0,
        )
        edited_permissions = st.multiselect("Aree autorizzate", PERMISSION_AREAS, default=selected_permissions)
        c3, c4 = st.columns(2)
        edited_admin = c3.checkbox("Amministratore", value=bool(selected["is_admin"]))
        edited_active = c4.checkbox("Utente attivo", value=bool(selected["active"]))
        new_password = st.text_input("Nuova password (lascia vuoto per non cambiarla)", type="password")
        save_user = st.form_submit_button("Salva modifiche", type="primary", use_container_width=True)
    if save_user:
        if not edited_name:
            st.error("Il nome visualizzato è obbligatorio.")
        elif new_password and len(new_password) < 8:
            st.error("La nuova password deve contenere almeno 8 caratteri.")
        elif not edited_admin and not edited_permissions:
            st.error("Seleziona almeno un’area autorizzata.")
        else:
            with db() as conn:
                conn.execute(
                    """
                    UPDATE app_users SET display_name = ?, operator_name = ?, permissions = ?,
                    active = ?, is_admin = ? WHERE username = ?
                    """,
                    (
                        clean(edited_name), upper(edited_operator), json.dumps(edited_permissions),
                        int(edited_active), int(edited_admin), selected_username,
                    ),
                )
                if new_password:
                    conn.execute(
                        "UPDATE app_users SET password_hash = ? WHERE username = ?",
                        (password_hash(new_password), selected_username),
                    )
            st.success("Utente aggiornato.")
            st.rerun()

    confirm_delete = st.checkbox(f"Confermo l’eliminazione dell’utente {selected_username}")
    if st.button("Elimina utente", disabled=not confirm_delete, use_container_width=True):
        with db() as conn:
            conn.execute("DELETE FROM app_users WHERE username = ?", (selected_username,))
        st.success("Utente eliminato.")
        st.rerun()


def login():
    configured = str(st.secrets.get("ADMIN_PASSWORD", "") or "").strip()
    if st.session_state.get("authenticated"):
        if "current_user" not in st.session_state:
            st.session_state["current_user"] = {
                "username": "admin", "display_name": "Amministratore", "operator_name": "",
                "permissions": PERMISSION_AREAS, "is_admin": True,
            }
        return True
    st.title("Gestionale Ancillary")
    username = st.text_input("Username", value="admin")
    password = st.text_input("Password", type="password")
    if st.button("Accedi", type="primary", use_container_width=True):
        normalized_username = clean(username).lower()
        if normalized_username == "admin" and configured and hmac.compare_digest(password, configured):
            st.session_state["authenticated"] = True
            st.session_state["current_user"] = {
                "username": "admin", "display_name": "Amministratore", "operator_name": "",
                "permissions": PERMISSION_AREAS, "is_admin": True,
            }
            st.rerun()
        with db() as conn:
            row = conn.execute(
                "SELECT * FROM app_users WHERE username = ? AND active = 1", (normalized_username,)
            ).fetchone()
        if row and password_matches(password, row["password_hash"]):
            st.session_state["authenticated"] = True
            st.session_state["current_user"] = {
                "username": row["username"], "display_name": row["display_name"],
                "operator_name": row["operator_name"],
                "permissions": json.loads(row["permissions"] or "[]"), "is_admin": bool(row["is_admin"]),
            }
            st.rerun()
        st.error("Username o password errati.")
    if not configured:
        with db() as conn:
            user_count = conn.execute("SELECT COUNT(*) FROM app_users WHERE active = 1").fetchone()[0]
        if user_count == 0:
            st.warning("Configura ADMIN_PASSWORD nei Secrets per attivare il primo accesso amministratore.")
    return False


init_db()
if st.query_params.get("ddt_sign"):
    signing_page(st.query_params.get("ddt_sign"))
    st.stop()
if not login():
    st.stop()

st.title("Gestionale Noleggi, Ancillary e Danni")
st.caption("Importazione automatica dei report, archivio contratti, statistiche, ancillary e addebito danni")
all_data = load_data()
all_contracts = load_contracts()
all_damages = load_damages()
all_cash_movements = load_cash_movements()
all_special_events = load_special_events()

navigation = {
    "ANCILLARY": [
        ("📊 Dashboard ancillary", "Dashboard ancillary"),
        ("📈 Analisi ancillary RA", "Analisi ancillary RA"),
        ("🗂️ Archivio ancillary", "Archivio ancillary"),
        ("➕ Inserimento ancillary", "Inserimento ancillary"),
    ],
    "CONTRATTI RA": [
        ("📊 Analisi contratti RA", "Analisi contratti RA"),
        ("🗂️ Archivio contratti RA", "Archivio contratti RA"),
        ("🔀 Analisi incrociata", "Analisi incrociata"),
    ],
    "ADDEBITO DANNI": [
        ("📱 Check-in rapido danni", "Check-in rapido danni"),
        ("📊 Analisi addebito danni", "Analisi addebito danni"),
        ("🗂️ Archivio addebito danni", "Archivio addebito danni"),
        ("➕ Inserimento addebito danni", "Inserimento addebito danni"),
    ],
    "CASSA": [
        ("📊 Dashboard cassa", "Dashboard cassa"),
        ("➕ Inserimento movimento", "Inserimento movimento cassa"),
        ("🗂️ Archivio movimenti", "Archivio movimenti cassa"),
        ("🧮 Conteggio e chiusura", "Conteggio e chiusura cassa"),
        ("📥 Importazione storico", "Importazione storico cassa"),
    ],
    "EVENTI SPECIALI": [
        ("📊 Dashboard eventi", "Dashboard eventi speciali"),
        ("➕ Crea evento", "Crea evento speciale"),
        ("📅 Calendario eventi", "Calendario eventi speciali"),
        ("🚗 Veicoli evento", "Veicoli eventi speciali"),
        ("👥 Partecipanti e personale", "Partecipanti eventi speciali"),
        ("📎 Documenti e fotografie", "Allegati eventi speciali"),
        ("🗂️ Archivio eventi", "Archivio eventi speciali"),
    ],
    "DOCUMENTI": [("📄 Documenti di trasporto", "Documenti di trasporto")],
    "AMMINISTRAZIONE": [
        ("🔐 Utenti e autorizzazioni", "Utenti e autorizzazioni"),
        ("👥 Configurazione operatori", "Configurazione operatori"),
        ("💾 Importazione e backup", "Importazione e backup"),
    ],
}

visible_navigation = {
    area: functions for area, functions in navigation.items() if area in allowed_areas()
}
if not user_is_admin() and "AMMINISTRAZIONE" in visible_navigation:
    visible_navigation["AMMINISTRAZIONE"] = [
        item for item in visible_navigation["AMMINISTRAZIONE"] if item[1] != "Utenti e autorizzazioni"
    ]

visible_pages = [page_name for functions in visible_navigation.values() for _, page_name in functions]
if not visible_pages:
    st.error("Questo utente non ha aree autorizzate. Contatta l’amministratore.")
    st.stop()
if st.session_state.get("navigation_page") not in visible_pages:
    st.session_state.navigation_page = visible_pages[0]

logged_user = current_user()
st.sidebar.success(f"👤 {logged_user.get('display_name') or logged_user.get('username')}")
if st.sidebar.button("Esci", use_container_width=True):
    for key in ["authenticated", "current_user", "navigation_page"]:
        st.session_state.pop(key, None)
    st.rerun()
st.sidebar.subheader("Menu")
for area, functions in visible_navigation.items():
    area_pages = [page_name for _, page_name in functions]
    with st.sidebar.expander(area, expanded=st.session_state.navigation_page in area_pages):
        for button_label, page_name in functions:
            if st.button(
                button_label,
                key=f"nav_{area}_{page_name}",
                type="primary" if st.session_state.navigation_page == page_name else "secondary",
                use_container_width=True,
            ):
                st.session_state.navigation_page = page_name
                st.rerun()

page = st.session_state.navigation_page

# I filtri generali ancillary sono mostrati solo nelle pagine che li usano.
if page in {"Dashboard ancillary", "Archivio ancillary"} and not all_data.empty:
    filtered = filters(all_data)
else:
    filtered = all_data

if page == "Dashboard ancillary":
    dashboard(filtered)
elif page == "Analisi contratti RA":
    contracts_dashboard(all_contracts)
elif page == "Analisi ancillary RA":
    ancillary_ra_dashboard(all_contracts)
elif page == "Analisi incrociata":
    combined_analysis(all_contracts, all_data)
elif page == "Analisi addebito danni":
    damage_dashboard(all_damages)
elif page == "Archivio addebito danni":
    damage_archive(all_damages)
elif page == "Inserimento addebito danni":
    damage_form(all_damages)
elif page == "Check-in rapido danni":
    quick_damage_checkin()
elif page == "Dashboard cassa":
    cash_dashboard(all_cash_movements)
elif page == "Inserimento movimento cassa":
    cash_movement_form(all_cash_movements)
elif page == "Archivio movimenti cassa":
    cash_archive(all_cash_movements)
elif page == "Conteggio e chiusura cassa":
    cash_closing_page(all_cash_movements)
elif page == "Importazione storico cassa":
    cash_import_page()
elif page == "Dashboard eventi speciali":
    special_events_dashboard(all_special_events)
elif page == "Crea evento speciale":
    special_event_form(all_special_events)
elif page == "Calendario eventi speciali":
    events_calendar(all_special_events)
elif page == "Veicoli eventi speciali":
    event_vehicles_page(all_special_events)
elif page == "Partecipanti eventi speciali":
    event_participants_page(all_special_events)
elif page == "Allegati eventi speciali":
    event_files_page(all_special_events)
elif page == "Archivio eventi speciali":
    events_archive(all_special_events)
elif page == "Archivio ancillary":
    archive(filtered)
elif page == "Archivio contratti RA":
    contracts_archive(all_contracts)
elif page == "Inserimento ancillary":
    record_form(all_data)
elif page == "Documenti di trasporto":
    documents_page()
elif page == "Configurazione operatori":
    operator_settings()
elif page == "Utenti e autorizzazioni":
    user_settings()
else:
    import_backup()
